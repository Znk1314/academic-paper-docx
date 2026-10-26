"""学术论文 / 研究报告 Word 生成库。

严格对齐《学术论文 Word 排版与生成规范》（references/format-spec.md）。
把文档当成"格式系统 + 内容模型"来生成：所有格式参数集中在 SPEC，改主题只改内容模型。

典型用法见 templates/report_template.py 或 build_docx.py（JSON 驱动）。

依赖：python-docx >= 1.1
"""

from __future__ import annotations

import copy
import os
import re

import palette
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

try:  # 作为脚本运行时同目录导入；作为包导入时走相对导入
    from omml_latex import add_latex_to_paragraph
except ImportError:  # pragma: no cover
    from .omml_latex import add_latex_to_paragraph

# ---------------------------------------------------------------- 规范参数

SPEC = {
    "page": {
        "width_cm": 21.0, "height_cm": 29.7,
        "top_cm": 2.4, "bottom_cm": 2.2, "left_cm": 2.6, "right_cm": 2.4,
        "header_cm": 1.2, "footer_cm": 1.0,
        "different_first_page": True,
    },
    # 文档版式色：全部来自 palette.py 的 DOC 角色（单一色值来源），
    # 彩色数 = 4（深蓝 / 蓝 / 橙 / 绿），彩色数与灰度可辨性由 qa_check 校验。
    "colors": {
        "navy": palette.DOC["heading"].lstrip("#"),     # 一级标题（中性深蓝）
        "blue": palette.DOC["accent"].lstrip("#"),      # 二级标题 / 强调线（Okabe-Ito blue）
        "orange": palette.DOC["warn"].lstrip("#"),      # 辅助强调（Okabe-Ito orange）
        "green": palette.DOC["ok"].lstrip("#"),         # 正向 / 完成（Okabe-Ito green）
        "ink": palette.DOC["ink"].lstrip("#"),          # 正文
        "gray": palette.DOC["muted"].lstrip("#"),       # 次级说明
        "pale": palette.DOC["pale"].lstrip("#"),        # 浅色提示框
        "pale_line": palette.DOC["rule"].lstrip("#"),   # 浅色边框
    },
    "fonts": {
        "head_cn": "黑体", "head_en": "Arial",
        "body_cn": "宋体", "body_en": "Times New Roman",
    },
    "sizes": {
        "cover_kicker": 15, "cover_title": 24, "cover_subtitle": 20,
        "cover_note": 11, "cover_info": 12, "cover_bottom": 12,
        "h1": 16, "h2": 13,
        "body": 12, "abstract_en": 11, "keywords": 10.5,
        "caption": 9.5, "table": 9, "note": 8.5,
        "header": 8.5, "footer": 8, "reference": 9,
        "toc_title": 18, "toc_item": 11,
    },
    "text_width_cm": 15.8,
    "table_width_cm": 15.5,
    # 插图尺寸（GB/T 7713.2—2022 5.4.1：大小适当、图中文字清晰可见）
    # 版心宽 = 21 − 2.6 − 2.4 = 16.0 cm；通栏取 15.5 cm 与三线表同宽，留安全余量。
    "figure_half_cm": 7.5,        # 半栏 75 mm
    "figure_full_cm": 15.5,       # 通栏 155 mm
    "figure_default_cm": 15.5,
    "figure_max_cm": 15.5,        # 不得超过版心
    "figure_max_height_cm": 22.0,  # 图 + 图题 + 图注须容于一页
    "figure_min_font_pt": 6.0,    # 插入后图内文字有效字号下限
    "cover_top_space_pt": 104,
    "cover_bottom_space_pt": 132,
    "table_header_color": "navy",   # 表头文字色；想更素净可改成 "ink"
    "toc_include_front_matter": False,  # True 时摘要/Abstract 也进目录
}

# ---------------------------------------------------------------- 元素顺序表
# OOXML 对子元素顺序敏感；用 insert_ordered 保证不产生"顺序非法"的 XML。

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
    "pPrChange",
]
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
    "tblDescription",
]
TRPR_ORDER = [
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
    "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
]
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]


def _local(el) -> str:
    tag = el.tag
    return tag.split("}")[-1] if isinstance(tag, str) else ""


def rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str.upper())


def insert_ordered(parent, el, order):
    """把 el 按 OOXML 子元素顺序插入 parent。"""
    tag = _local(el)
    if tag not in order:
        parent.append(el)
        return el
    idx = order.index(tag)
    for child in list(parent):
        ctag = _local(child)
        if ctag in order and order.index(ctag) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def remove_children(parent, tag: str):
    for el in parent.findall(qn("w:" + tag)):
        parent.remove(el)


def deep_merge(base: dict, override: dict) -> dict:
    out = copy.deepcopy(base)
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = deep_merge(out[k], v)
        else:
            out[k] = v
    return out


# ---------------------------------------------------------------- run / 段落

def set_run_font(run, cn="宋体", en="Times New Roman", size=12,
                 bold=False, color=None, italic=False):
    """设置中西文双字体 + 字号 + 粗细 + 颜色（去掉主题字体，保证 Word 里生效）。"""
    run.font.name = en
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color if isinstance(color, RGBColor) else rgb(color)
    rPr = run._element.get_or_add_rPr()
    _force_rfonts(rPr, cn, en)
    return run


def _force_rfonts(rPr, cn, en):
    """在 rPr 中写入 w:rFonts（放在 rStyle 之后，符合 OOXML 顺序）。"""
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rStyle = rPr.find(qn("w:rStyle"))
        if rStyle is not None:
            rStyle.addnext(rFonts)
        else:
            rPr.insert(0, rFonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        if rFonts.get(qn("w:" + attr)) is not None:
            del rFonts.attrib[qn("w:" + attr)]
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)
    return rFonts


def set_style_font(style, cn, en, size, bold=False, color=None, italic=False):
    style.font.name = en
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color if isinstance(color, RGBColor) else rgb(color)
    rPr = style.element.get_or_add_rPr()
    _force_rfonts(rPr, cn, en)


def set_first_line_chars(pf, chars):
    """按"字符"设置首行缩进（Word 的 2 字符缩进，比固定磅值更稳）。"""
    pPr = pf._element.get_or_add_pPr() if hasattr(pf, "_element") else pf
    ind = pPr.get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(int(chars)))


def _to_length(value):
    """把"磅值 / Length / EMU 整数"统一成 Length。

    注意：`-Cm(0.74)` 在 Python 里会退化成普通 int（因为 Length 是 int 子类，
    取负后丢掉类型），所以这里按量级判断：> 20000 视为 EMU，否则视为磅。
    """
    if hasattr(value, "twips"):
        return value
    number = float(value)
    if abs(number) > 20000:
        return Emu(int(round(number)))
    return Pt(number)


def para_format(p, *, align=None, line_spacing=None, space_before=None,
                space_after=None, first_line_indent=None, first_line_chars=None,
                left_indent=None, right_indent=None, keep_with_next=None,
                keep_together=None, page_break_before=None, style=None,
                contextual_spacing=None, hanging_cm=None):
    if style is not None:
        p.style = style
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if left_indent is not None:
        pf.left_indent = left_indent if hasattr(left_indent, "twips") else Cm(left_indent)
    if right_indent is not None:
        pf.right_indent = right_indent if hasattr(right_indent, "twips") else Cm(right_indent)
    if hanging_cm is not None:
        pPr = p._p.get_or_add_pPr()
        ind = pPr.get_or_add_ind()
        twips = str(int(round(hanging_cm * 567)))
        if ind.get(qn("w:left")) is None:
            ind.set(qn("w:left"), twips)
        ind.set(qn("w:firstLine"), "0")
        ind.set(qn("w:hanging"), twips)
        first_line_chars = None
    elif first_line_indent is not None:
        value = _to_length(first_line_indent)
        if value < 0:
            pPr = p._p.get_or_add_pPr()
            ind = pPr.get_or_add_ind()
            twips = str(abs(int(round(value.twips))))
            ind.set(qn("w:firstLine"), "0")
            ind.set(qn("w:hanging"), twips)
            first_line_chars = None
        else:
            pf.first_line_indent = value
            if first_line_chars is None:
                first_line_chars = 0
    if first_line_chars is not None:
        set_first_line_chars(pf, first_line_chars)
    if keep_with_next is not None:
        pf.keep_with_next = keep_with_next
    if keep_together is not None:
        pf.keep_together = keep_together
    if page_break_before is not None:
        pf.page_break_before = page_break_before
    if contextual_spacing is not None:
        pPr = p._p.get_or_add_pPr()
        cs = OxmlElement("w:contextualSpacing")
        if not contextual_spacing:
            cs.set(qn("w:val"), "0")
        insert_ordered(pPr, cs, PPR_ORDER)
    return p


def add_paragraph_border(p_or_style, edge="bottom", sz=6, color="0072B2", space=4):
    """给段落/样式加一条边线。sz 单位是 1/8 磅（6 = 0.75 pt，12 = 1.5 pt）。"""
    if hasattr(p_or_style, "_p"):
        pPr = p_or_style._p.get_or_add_pPr()
    elif hasattr(p_or_style, "element"):
        pPr = p_or_style.element.get_or_add_pPr()
    else:
        pPr = p_or_style
    remove_children(pPr, "pBdr")
    pBdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)
    insert_ordered(pPr, pBdr, PPR_ORDER)
    return pBdr


def shade_paragraph(p, fill="EAF2FA", border_color="B4C6E7", sz=6):
    pPr = p._p.get_or_add_pPr()
    remove_children(pPr, "shd")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(pPr, shd, PPR_ORDER)

    remove_children(pPr, "pBdr")
    pBdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), border_color)
        pBdr.append(el)
    insert_ordered(pPr, pBdr, PPR_ORDER)


def _field_part(paragraph, fld_type, instr=None, text=None, font=None):
    """在段落里追加一个域片段（begin / instrText / separate / text / end）。
    每个片段独占一个 run，这是 Word 最稳的域结构（避免目录变成原始域代码）。"""
    run = paragraph.add_run()
    if font:
        set_run_font(run, *font)
    if fld_type == "instr":
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = instr
    elif fld_type == "text":
        el = OxmlElement("w:t")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    else:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), fld_type)
        if fld_type == "begin":
            el.set(qn("w:dirty"), "true")
    run._r.append(el)
    return run


def add_field(paragraph, instr: str, placeholder: str = "", size=None,
              cn="宋体", en="Times New Roman", color=None, bold=False):
    """插入一个 Word 域（PAGE / NUMPAGES 等）。"""
    font = None
    if size is not None:
        font = (cn, en, size, bold,
                color if isinstance(color, RGBColor) or color is None
                else rgb(color))
    _field_part(paragraph, "begin", font=font)
    _field_part(paragraph, "instr", instr=" %s " % instr.strip(), font=font)
    if placeholder:
        _field_part(paragraph, "separate", font=font)
        _field_part(paragraph, "text", text=placeholder, font=font)
    _field_part(paragraph, "end", font=font)


# ---------------------------------------------------------------- 三线表

def set_cell_borders(cell, **edges):
    """edges: top=dict(sz=..,color=..) / bottom=.. ；未指定的边设为无。"""
    tcPr = cell._tc.get_or_add_tcPr()
    remove_children(tcPr, "tcBorders")
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        spec = edges.get(edge)
        if spec:
            el.set(qn("w:val"), spec.get("val", "single"))
            el.set(qn("w:sz"), str(spec.get("sz", 8)))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec.get("color", "000000"))
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    insert_ordered(tcPr, borders, TCPR_ORDER)


def set_three_line_table(table, outer_sz=12, header_sz=6, color="000000"):
    """三线表：顶/底线 1.5 pt(sz=12)、栏目线 0.75 pt(sz=6)、无竖线。

    GB/T 7713.2 规定的是三线表的**构成**（顶线/栏目线/底线，不画竖线）；
    具体磅值属期刊惯例，顶底 1.5 磅、栏目线 0.5—0.75 磅是通行区间，这里取 0.75。
    """
    tbl = table._tbl
    tblPr = tbl.tblPr
    remove_children(tblPr, "tblBorders")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        if edge in ("top", "bottom"):
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(outer_sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    insert_ordered(tblPr, borders, TBLPR_ORDER)

    for cell in table.rows[0].cells:
        set_cell_borders(cell, bottom={"sz": header_sz, "color": color})


def set_table_repeat_header(table):
    trPr = table.rows[0]._tr.get_or_add_trPr()
    remove_children(trPr, "tblHeader")
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    insert_ordered(trPr, tblHeader, TRPR_ORDER)


def forbid_row_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        remove_children(trPr, "cantSplit")
        insert_ordered(trPr, OxmlElement("w:cantSplit"), TRPR_ORDER)


def set_table_layout_fixed(table, width_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    remove_children(tblPr, "tblW")
    remove_children(tblPr, "tblLayout")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(width_cm * 567)))  # 1 cm = 567 twips
    insert_ordered(tblPr, tblW, TBLPR_ORDER)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    insert_ordered(tblPr, layout, TBLPR_ORDER)
    # 单元格边距：上下 0，左右 0.15 cm，避免行高被默认边距撑大
    remove_children(tblPr, "tblCellMar")
    mar = OxmlElement("w:tblCellMar")
    for edge, twips in (("top", 0), ("left", 85), ("bottom", 0), ("right", 85)):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    insert_ordered(tblPr, mar, TBLPR_ORDER)


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    remove_children(tblPr, "tblBorders")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "nil")
        borders.append(el)
    insert_ordered(tblPr, borders, TBLPR_ORDER)


def set_col_widths(table, widths_cm):
    for i, w in enumerate(widths_cm):
        for row in table.rows:
            if i < len(row.cells):
                row.cells[i].width = Cm(w)
    # 同时修正 tblGrid，保证 Word 按固定列宽排版
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            if i < len(widths_cm):
                gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))


# ---------------------------------------------------------------- 页面 / 样式

def setup_page(doc, spec):
    page = spec["page"]
    sec = doc.sections[0]
    sec.page_width = Cm(page["width_cm"])
    sec.page_height = Cm(page["height_cm"])
    sec.top_margin = Cm(page["top_cm"])
    sec.bottom_margin = Cm(page["bottom_cm"])
    sec.left_margin = Cm(page["left_cm"])
    sec.right_margin = Cm(page["right_cm"])
    sec.header_distance = Cm(page["header_cm"])
    sec.footer_distance = Cm(page["footer_cm"])
    sec.gutter = Cm(0)
    sec.different_first_page_header_footer = bool(page["different_first_page"])
    return sec


def _setup_doc_defaults(doc, spec):
    """docDefaults：主题字体不参与，语言设为中文，避免中文字体回退。"""
    styles_el = doc.styles.element
    rpr_default = styles_el.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault"))
    if rpr_default is None:
        dd = styles_el.find(qn("w:docDefaults"))
        if dd is None:
            dd = OxmlElement("w:docDefaults")
            styles_el.insert(0, dd)
        rpr_default = OxmlElement("w:rPrDefault")
        dd.append(rpr_default)
    rPr = rpr_default.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        rpr_default.append(rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        if rFonts.get(qn("w:" + attr)) is not None:
            del rFonts.attrib[qn("w:" + attr)]
    rFonts.set(qn("w:ascii"), spec["fonts"]["body_en"])
    rFonts.set(qn("w:hAnsi"), spec["fonts"]["body_en"])
    rFonts.set(qn("w:eastAsia"), spec["fonts"]["body_cn"])
    rFonts.set(qn("w:cs"), spec["fonts"]["body_en"])
    remove_children(rPr, "lang")
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    rPr.append(lang)


def setup_styles(doc, spec):
    sizes, colors, fonts = spec["sizes"], spec["colors"], spec["fonts"]

    # --- 正文
    normal = doc.styles["Normal"]
    set_style_font(normal, fonts["body_cn"], fonts["body_en"],
                   sizes["body"], False, rgb(colors["ink"]))
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(24)   # 12 pt 正文的 2 字符；逐段再用 firstLineChars 精确控制
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    # --- 一级标题
    h1 = doc.styles["Heading 1"]
    set_style_font(h1, fonts["head_cn"], fonts["head_en"],
                   sizes["h1"], True, rgb(colors["navy"]))
    pf = h1.paragraph_format
    pf.space_before = Pt(16)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.25
    pf.first_line_indent = Pt(0)
    set_first_line_chars(pf, 0)
    pf.left_indent = Cm(0)
    pf.keep_with_next = True
    pf.keep_together = True
    pf.page_break_before = False
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_paragraph_border(h1, "bottom", sz=6, color=colors["blue"], space=4)
    _set_outline_level(h1.element.get_or_add_pPr(), 0)

    # --- 二级标题
    h2 = doc.styles["Heading 2"]
    set_style_font(h2, fonts["head_cn"], fonts["head_en"],
                   sizes["h2"], True, rgb(colors["blue"]))
    pf = h2.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.25
    pf.first_line_indent = Pt(0)
    set_first_line_chars(pf, 0)
    pf.left_indent = Cm(0)
    pf.keep_with_next = True
    pf.keep_together = True
    pf.page_break_before = False
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_outline_level(h2.element.get_or_add_pPr(), 1)

    # --- 前置部分标题（摘要 / Abstract）：外观同一级标题，但不进目录
    _ensure_front_heading_style(doc, spec)

    # --- 目录条目样式（Word 生成目录后套用 TOC1）
    _ensure_toc_style(doc, spec)

    # --- 页眉页脚
    for name in ("Header", "Footer"):
        st = doc.styles[name]
        set_style_font(st, fonts["body_cn"], fonts["body_en"],
                       sizes["header"] if name == "Header" else sizes["footer"],
                       False, rgb(colors["gray"]))
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.first_line_indent = Pt(0)
        set_first_line_chars(st.paragraph_format, 0)


def _set_outline_level(pPr, level: int):
    remove_children(pPr, "outlineLvl")
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(level))
    insert_ordered(pPr, el, PPR_ORDER)


def _ensure_front_heading_style(doc, spec):
    """摘要/Abstract 用的"像一级标题但不进目录"的样式（大纲级别 9 = 正文）。"""
    name = "Front Heading"
    existing = [s for s in doc.styles if s.name == name]
    if existing:
        style = existing[0]
    else:
        from docx.enum.style import WD_STYLE_TYPE
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Heading 1"]
    set_style_font(style, spec["fonts"]["head_cn"], spec["fonts"]["head_en"],
                   spec["sizes"]["h1"], True, rgb(spec["colors"]["navy"]))
    pf = style.paragraph_format
    pf.space_before = Pt(16)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.25
    pf.first_line_indent = Pt(0)
    set_first_line_chars(pf, 0)
    pf.keep_with_next = True
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_paragraph_border(style, "bottom", sz=6,
                         color=spec["colors"]["blue"], space=4)
    _set_outline_level(style.element.get_or_add_pPr(), 9)
    return style


def _ensure_toc_style(doc, spec):
    """注入 Word 内置 TOC1 样式（styleId=TOC1），保证更新目录后字体受控。"""
    styles_el = doc.styles.element
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == "TOC1":
            styles_el.remove(st)
    sz = int(spec["sizes"]["toc_item"] * 2)  # 半磅
    xml = (
        '<w:style xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' w:type="paragraph" w:styleId="TOC1">'
        '<w:name w:val="toc 1"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/>'
        '<w:uiPriority w:val="39"/><w:unhideWhenUsed/>'
        '<w:pPr><w:spacing w:before="0" w:after="60" w:line="300" w:lineRule="auto"/>'
        '<w:ind w:left="0" w:firstLine="0"/><w:jc w:val="left"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"'
        ' w:eastAsia="宋体"/><w:color w:val="%s"/><w:sz w:val="%d"/></w:rPr>'
        '</w:style>' % (spec["colors"]["ink"], sz)
    )
    from lxml import etree
    styles_el.append(etree.fromstring(xml))


SETTINGS_ORDER = [
    "writeProtection", "view", "zoom", "removePersonalInformation",
    "removeDateAndTime", "doNotDisplayPageBoundaries", "displayBackgroundShape",
    "printPostScriptOverText", "printFractionalCharacterWidth", "printFormsData",
    "embedTrueTypeFonts", "embedSystemFonts", "saveSubsetFonts", "saveFormsData",
    "mirrorMargins", "alignBordersAndEdges", "bordersDoNotSurroundHeader",
    "bordersDoNotSurroundFooter", "gutterAtTop", "hideSpellingErrors",
    "hideGrammaticalErrors", "activeWritingStyle", "proofState", "formsDesign",
    "attachedTemplate", "linkStyles", "stylePaneFormatFilter",
    "stylePaneSortMethod", "documentType", "mailMerge", "revisionView",
    "trackRevisions", "doNotTrackMoves", "doNotTrackFormatting",
    "documentProtection", "autoFormatOverride", "styleLockTheme",
    "styleLockQFSet", "defaultTabStop", "autoHyphenation",
    "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps",
    "showEnvelope", "summaryLength", "clickAndTypeStyle", "defaultTableStyle",
    "evenAndOddHeaders", "bookFoldRevPrinting", "bookFoldPrinting",
    "bookFoldPrintingSheets", "drawingGridHorizontalSpacing",
    "drawingGridVerticalSpacing", "displayHorizontalDrawingGridEvery",
    "displayVerticalDrawingGridEvery", "doNotUseMarginsForDrawingGridOrigin",
    "drawingGridHorizontalOrigin", "drawingGridVerticalOrigin",
    "doNotShadeFormData", "noPunctuationKerning", "characterSpacingControl",
    "printTwoOnOne", "strictFirstAndLastChars", "noLineBreaksAfter",
    "noLineBreaksBefore", "savePreviewPicture", "doNotValidateAgainstSchema",
    "saveInvalidXml", "ignoreMixedContent", "alwaysShowPlaceholderText",
    "doNotDemarcateInvalidXml", "saveXmlDataOnly", "useXSLTWhenSaving",
    "saveThroughXslt", "showXMLTags", "alwaysMergeEmptyNamespace",
    "updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr", "compat",
    "docVars", "rsids", "mathPr", "attachedSchema", "themeFontLang",
    "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures",
    "forceUpgrade", "captions", "readModeInkLockDown", "schemaLibrary",
    "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator",
]


def _set_setting(doc, tag, attrs):
    settings = doc.settings.element
    remove_children(settings, tag)
    el = OxmlElement("w:" + tag)
    for k, v in attrs.items():
        el.set(qn("w:" + k), str(v))
    insert_ordered(settings, el, SETTINGS_ORDER)
    return el


def enable_update_fields(doc):
    """打开文档时提示/自动更新域（配合 TOC dirty）。"""
    _set_setting(doc, "updateFields", {"val": "true"})


def set_print_view(doc, zoom=100):
    """默认以页面视图 / 打印布局打开，缩放 100%（spec §18）。"""
    _set_setting(doc, "view", {"val": "print"})
    _set_setting(doc, "zoom", {"percent": zoom})


def apply_print_view_to_file(path, zoom=100):
    """最终保存后调用：把已存在的 docx 重新设为页面视图（不经过 Word）。"""
    doc = Document(path)
    set_print_view(doc, zoom)
    doc.save(path)


# ---------------------------------------------------------------- 构建器

class PaperBuilder:
    def __init__(self, spec_overrides=None, base_dir=None):
        self.spec = deep_merge(SPEC, spec_overrides or {})
        self.base_dir = os.path.abspath(base_dir) if base_dir else None
        self.doc = None
        self.tbl_no = 0
        self.fig_no = 0
        self.eq_no = 0
        self.h1_no = 0
        self.warnings = []
        self._body_started = False

    # ---------- 通用小工具
    def _resolve_path(self, path):
        """相对路径先按 base_dir（通常是内容模型所在目录）解析，再退回当前目录。"""
        if not path:
            return path
        if os.path.isabs(path):
            return path
        candidates = []
        if self.base_dir:
            candidates.append(os.path.join(self.base_dir, path))
        candidates.append(os.path.abspath(path))
        for cand in candidates:
            if os.path.exists(cand):
                return os.path.abspath(cand)
        return os.path.abspath(candidates[0])

    def _c(self, key):
        return self.spec["colors"][key]

    def _s(self, key):
        return self.spec["sizes"][key]

    def _f(self, key):
        return self.spec["fonts"][key]

    def _para(self, align=None, **kw):
        p = self.doc.add_paragraph()
        para_format(p, **kw)
        if align is not None:
            p.alignment = align
        return p

    def _run(self, p, text, cn=None, en=None, size=None, bold=False,
             color=None, italic=False):
        r = p.add_run(text)
        set_run_font(r, cn or self._f("body_cn"), en or self._f("body_en"),
                     size if size is not None else self._s("body"),
                     bold, rgb(color) if isinstance(color, str) else color,
                     italic)
        return r

    # ---------- 页面 / 基础
    def setup(self):
        self.doc = Document()
        setup_page(self.doc, self.spec)
        _setup_doc_defaults(self.doc, self.spec)
        setup_styles(self.doc, self.spec)
        props = self.doc.core_properties
        props.author = ""
        props.last_modified_by = ""
        props.comments = ""
        return self.doc

    def page_break(self):
        p = self.doc.add_paragraph()
        para_format(p, first_line_indent=0, first_line_chars=0,
                    space_before=0, space_after=0, line_spacing=1.0)
        p.add_run().add_break(WD_BREAK.PAGE)
        return p

    # ---------- 封面
    def add_cover(self, meta):
        doc = self.doc
        S, C = self.spec["sizes"], self.spec["colors"]

        spacer = doc.add_paragraph()
        para_format(spacer, first_line_indent=0, first_line_chars=0,
                    space_before=self.spec["cover_top_space_pt"],
                    space_after=0, line_spacing=1.0)
        set_run_font(spacer.add_run(""), "宋体", "Times New Roman", 6)

        kicker = " ".join(x for x in [meta.get("org", ""), meta.get("doc_type", "")] if x)
        if kicker:
            p = doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                        first_line_chars=0, space_before=0, space_after=26,
                        line_spacing=1.2)
            self._run(p, kicker, self._f("head_cn"), self._f("head_en"),
                      S["cover_kicker"], False, C["blue"])

        title = meta.get("title", "")
        if title:
            p = doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                        first_line_chars=0, space_before=0, space_after=10,
                        line_spacing=1.25)
            self._run(p, title, self._f("head_cn"), self._f("head_en"),
                      S["cover_title"], True, C["navy"])

        subtitle = meta.get("subtitle", "")
        if subtitle:
            p = doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                        first_line_chars=0, space_before=0, space_after=8,
                        line_spacing=1.2)
            self._run(p, subtitle, self._f("head_cn"), self._f("head_en"),
                      S["cover_subtitle"], False, C["blue"])

        note = meta.get("note", "")
        if note:
            p = doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                        first_line_chars=0, space_before=0, space_after=6,
                        line_spacing=1.2)
            self._run(p, note, self._f("body_cn"), self._f("body_en"),
                      S["cover_note"], False, C["gray"])

        # 学术蓝横线 1.5 pt
        rule = doc.add_paragraph()
        para_format(rule, first_line_indent=0, first_line_chars=0,
                    space_before=14, space_after=20, line_spacing=1.0)
        add_paragraph_border(rule, "bottom", sz=12, color=C["blue"], space=1)
        set_run_font(rule.add_run(""), "宋体", "Times New Roman", 2)

        info = meta.get("info") or []
        if info:
            table = doc.add_table(rows=0, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_layout_fixed(table, 14.6)
            remove_table_borders(table)
            set_col_widths(table, [4.0, 10.6])
            for label, value in info:
                cells = table.add_row().cells
                cells[0].width = Cm(4.0)
                cells[1].width = Cm(10.6)
                for idx, (text, align, bold, color, cn) in enumerate([
                    (label, WD_ALIGN_PARAGRAPH.RIGHT, True, C["navy"], self._f("head_cn")),
                    (str(value), WD_ALIGN_PARAGRAPH.LEFT, False, C["ink"], self._f("body_cn")),
                ]):
                    cell = cells[idx]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    p = cell.paragraphs[0]
                    para_format(p, align=align, first_line_indent=0,
                                first_line_chars=0, space_before=2,
                                space_after=2, line_spacing=1.3)
                    self._run(p, text, cn, self._f("head_en") if bold else self._f("body_en"),
                              S["cover_info"], bold, color)

        bottom = doc.add_paragraph()
        para_format(bottom, first_line_indent=0, first_line_chars=0,
                    space_before=self.spec["cover_bottom_space_pt"],
                    space_after=0, line_spacing=1.0)
        set_run_font(bottom.add_run(""), "宋体", "Times New Roman", 6)

        org_line = " ".join(x for x in [meta.get("org", ""),
                                        meta.get("college", "")] if x)
        if org_line:
            p = doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                        first_line_chars=0, space_before=0, space_after=2,
                        line_spacing=1.2)
            self._run(p, org_line, self._f("head_cn"), self._f("head_en"),
                      S["cover_bottom"], False, C["gray"])
        date = meta.get("date", "")
        if date:
            p = doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                        first_line_chars=0, space_before=0, space_after=0,
                        line_spacing=1.2)
            self._run(p, date, self._f("head_cn"), self._f("head_en"),
                      S["cover_bottom"], False, C["gray"])

        if meta.get("title"):
            self.doc.core_properties.title = meta["title"]
            self.doc.core_properties.subject = meta.get("doc_type", "")

    # ---------- 标题
    def add_heading1(self, text, page_break=False, front=False):
        use_front = front and not self.spec.get("toc_include_front_matter")
        style = "Front Heading" if use_front else "Heading 1"
        p = self.doc.add_paragraph(style=style)
        para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=0,
                    first_line_chars=0, page_break_before=bool(page_break))
        self.h1_no += 1
        self._run(p, text, self._f("head_cn"), self._f("head_en"),
                  self._s("h1"), True, self._c("navy"))
        return p

    def add_heading2(self, text):
        p = self.doc.add_paragraph(style="Heading 2")
        para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=0,
                    first_line_chars=0)
        self._run(p, text, self._f("head_cn"), self._f("head_en"),
                  self._s("h2"), True, self._c("blue"))
        return p

    # ---------- 正文
    def add_body(self, text, indent_chars=2, size=None, align=None,
                 space_after=None, line_spacing=None):
        p = self.doc.add_paragraph()
        para_format(p, align=align or WD_ALIGN_PARAGRAPH.JUSTIFY,
                    line_spacing=line_spacing or 1.5,
                    space_before=0,
                    space_after=space_after if space_after is not None else 6,
                    first_line_indent=Pt(24) if indent_chars else Pt(0),
                    first_line_chars=indent_chars * 100)
        self._run(p, text, size=size if size is not None else self._s("body"),
                  color=self._c("ink"))
        return p

    def add_plain(self, text, size=None, color=None, align=None, indent_chars=0,
                  space_before=0, space_after=6, line_spacing=1.5,
                  cn=None, en=None, bold=False, italic=False,
                  keep_with_next=False):
        p = self.doc.add_paragraph()
        para_format(p, align=align or WD_ALIGN_PARAGRAPH.JUSTIFY,
                    line_spacing=line_spacing, space_before=space_before,
                    space_after=space_after,
                    first_line_indent=Pt(0) if not indent_chars else Pt(24),
                    first_line_chars=indent_chars * 100,
                    keep_with_next=keep_with_next)
        self._run(p, text, cn=cn, en=en,
                  size=size if size is not None else self._s("body"),
                  color=color if color is not None else self._c("ink"),
                  bold=bold, italic=italic)
        return p

    def add_list(self, items, start=1, size=None):
        """（1）（2）（3）编号段落，不用 Word 项目符号（spec §13）。"""
        for i, item in enumerate(items, start):
            p = self.doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                        space_before=2, space_after=2,
                        first_line_indent=Pt(0), first_line_chars=0)
            self._run(p, "（%d）%s" % (i, item),
                      size=size if size is not None else self._s("body"),
                      color=self._c("ink"))

    def add_callout(self, text, title=None):
        p = self.doc.add_paragraph()
        para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.4,
                    space_before=6, space_after=8,
                    first_line_indent=Pt(0), first_line_chars=0)
        shade_paragraph(p, fill=self._c("pale"), border_color=self._c("pale_line"))
        if title:
            self._run(p, title + "  ", size=11, bold=True, color=self._c("navy"))
        self._run(p, text, size=11, color=self._c("ink"))
        return p

    # ---------- 图表编号
    @staticmethod
    def _strip_number(text):
        return re.sub(r"^(图|表)\s*\d+(?:[-.]\d+)?\s*", "", (text or "").strip())

    def add_table_caption(self, caption, caption_en=None):
        self.tbl_no += 1
        body = self._strip_number(caption)
        text = "表%d　%s" % (self.tbl_no, body) if body else "表%d" % self.tbl_no
        p = self.doc.add_paragraph()
        para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                    first_line_chars=0, space_before=6, space_after=3,
                    line_spacing=1.2, keep_with_next=True)
        self._run(p, text, self._f("head_cn"), self._f("head_en"),
                  self._s("caption"), True, self._c("navy"))
        if caption_en:
            self._run(p, "  " + str(caption_en).strip(), self._f("head_cn"),
                      self._f("head_en"), self._s("caption"), False,
                      self._c("navy"))
        return p, text

    def add_figure_caption(self, caption, keep_with_next=False, caption_en=None):
        self.fig_no += 1
        body = self._strip_number(caption)
        text = "图%d　%s" % (self.fig_no, body) if body else "图%d" % self.fig_no
        p = self.doc.add_paragraph()
        para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                    first_line_chars=0, space_before=4, space_after=3,
                    line_spacing=1.2, keep_with_next=keep_with_next)
        self._run(p, text, self._f("head_cn"), self._f("head_en"),
                  self._s("caption"), True, self._c("navy"))
        if caption_en:
            self._run(p, "  " + str(caption_en).strip(), self._f("head_cn"),
                      self._f("head_en"), self._s("caption"), False,
                      self._c("navy"))
        return p, text

    def add_note(self, note, align=WD_ALIGN_PARAGRAPH.LEFT):
        text = note.strip()
        if not text.startswith("注"):
            text = "注：" + text
        # GB/T 7713.2—2022 5.4.3 / 5.5.3：图注、表注的末尾应加“。”
        if not text.endswith(("。", ".", "）", ")")):
            text = text + "。"
        p = self.doc.add_paragraph()
        para_format(p, align=align, first_line_indent=0, first_line_chars=0,
                    space_before=2, space_after=6, line_spacing=1.2)
        add_rich_text(p, text, self._f("body_cn"), self._f("body_en"),
                      self._s("note"), False, rgb(self._c("gray")))
        return p

    # ---------- 三线表
    def add_table(self, headers=None, rows=None, caption=None, note=None,
                  widths=None, aligns=None, table_size=None, columns=None,
                  caption_en=None, digits=None):
        """三线表。

        `columns` 给出列元信息（label/unit/digits/align），表头自动写成
        "量名称/单位"，数值按 `digits` 统一精度（规范 3.3：同一列小数位一致）。
        单元格里的 `^{a}` 会渲染为上标，用于显著性/脚标。
        """
        rows = rows or []
        if columns:
            headers = build_headers(columns)
            if digits is None:
                digits = [c.get("digits") for c in columns]
            if aligns is None:
                default_align = []
                for col in columns:
                    kind = str(col.get("align", "")).lower()
                    if kind in _ALIGN_MAP:
                        default_align.append(_ALIGN_MAP[kind])
                    elif kind == "decimal":
                        default_align.append(WD_ALIGN_PARAGRAPH.RIGHT)
                    elif col.get("digits") is not None:
                        default_align.append(WD_ALIGN_PARAGRAPH.RIGHT)
                    else:
                        default_align.append(None)
                aligns = default_align
        headers = headers or []

        total = self.spec["table_width_cm"]
        if caption:
            self.add_table_caption(caption, caption_en)
        ncol = len(headers)
        if not widths:
            widths = [total / ncol] * ncol
        scale = total / sum(widths)
        widths = [w * scale for w in widths]
        base_aligns = [WD_ALIGN_PARAGRAPH.CENTER] + [WD_ALIGN_PARAGRAPH.LEFT] * (ncol - 1)
        if aligns is None:
            aligns = base_aligns
        else:
            aligns = [x if not isinstance(x, str)
                      else _ALIGN_MAP.get(x.lower(), WD_ALIGN_PARAGRAPH.LEFT)
                      for x in aligns]
            aligns = [a if a is not None else (base_aligns[i] if i < ncol
                                               else WD_ALIGN_PARAGRAPH.LEFT)
                      for i, a in enumerate(aligns)]

        table = self.doc.add_table(rows=1, cols=ncol)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_layout_fixed(table, total)
        set_col_widths(table, widths)
        size = table_size or self._s("table")

        # 表头
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                        first_line_chars=0, space_before=1.5,
                        space_after=1.5, line_spacing=1.15)
            self._run(p, str(h), self._f("head_cn"), self._f("head_en"),
                      size, True,
                      self._c(self.spec.get("table_header_color", "navy")))

        # 数据行
        for row_data in rows:
            cells = table.add_row().cells
            for i, value in enumerate(row_data):
                cell = cells[i]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                para_format(p, align=aligns[i] if i < len(aligns) else WD_ALIGN_PARAGRAPH.LEFT,
                            first_line_indent=0, first_line_chars=0,
                            space_before=1.5, space_after=1.5, line_spacing=1.15)
                col_digits = None
                if digits is not None and i < len(digits):
                    col_digits = digits[i]
                text = format_number(value, col_digits)
                add_rich_text(p, text, self._f("body_cn"), self._f("body_en"),
                              size, False, rgb(self._c("ink")))

        set_three_line_table(table)
        set_table_repeat_header(table)
        forbid_row_split(table)
        if note:
            self.add_note(note)
        return table

    # ---------- 图片
    def add_figure(self, path, caption=None, note=None, width_cm=None,
                   caption_en=None):
        """插入插图并校验尺寸（GB/T 7713.2—2022 5.4.1）。

        三道约束：
        1. 宽度不得超过版心（默认 15.5 cm），超了自动收窄并记警告；
        2. 图 + 图题 + 图注须容于一页（默认图高 ≤22 cm）；
        3. 插入后图内文字有效字号 ≥6 pt —— 由 PNG 里记录的"设计宽度/基准字号"
           换算，缩得太小会打印不清（正是 5.4.1 说的"图中文字清晰可见"）。
        """
        path = self._resolve_path(path)
        if not os.path.exists(path):
            raise FileNotFoundError("图片不存在：%s" % path)
        name = os.path.basename(path)
        width_cm = float(width_cm or self.spec["figure_default_cm"])
        max_w = self.spec["figure_max_cm"]
        if width_cm > max_w:
            self.warnings.append("图片 %s 宽度 %.1f cm 超过版心 %.1f cm，已自动收窄"
                                 % (name, width_cm, max_w))
            width_cm = max_w

        # 原生尺寸 → 高度与缩放比
        try:
            from PIL import Image
            with Image.open(path) as img:
                nat_w, nat_h = img.size
                info = dict(img.info)
        except Exception:      # 缺 Pillow 或非 PIL 可读格式：跳过尺寸校验
            nat_w = nat_h = 0
            info = {}

        if nat_w and nat_h:
            height_cm = width_cm * nat_h / nat_w
            limit_h = self.spec["figure_max_height_cm"]
            if height_cm > limit_h:
                self.warnings.append(
                    "图片 %s 插入后高 %.1f cm 超过 %.1f cm，图+图题+图注可能跨页"
                    "（GB/T 7713.2—2022 5.4.5 要求插图紧随首次提及处、不拆散）"
                    % (name, height_cm, limit_h))

            software = info.get("Software") or ""
            design_w = design_pt = None
            for token in software.split():
                if token.startswith("width_mm="):
                    design_w = float(token.split("=")[1])
                elif token.startswith("font_pt="):
                    design_pt = float(token.split("=")[1])
            if design_w and design_pt:
                scale = width_cm * 10.0 / design_w
                effective = design_pt * scale
                floor = self.spec["figure_min_font_pt"]
                if effective < floor:
                    self.warnings.append(
                        "图片 %s 插入后图内文字仅约 %.1f pt（原设计 %.1f pt、"
                        "宽 %.0f mm，现插入 %.0f mm），低于 %.1f pt 印刷难辨"
                        "（GB/T 7713.2—2022 5.4.1）"
                        % (name, effective, design_pt, design_w,
                           width_cm * 10, floor))

        p = self.doc.add_paragraph()
        para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                    first_line_chars=0, space_before=6, space_after=2,
                    line_spacing=1.0, keep_with_next=True)
        run = p.add_run()
        run.add_picture(path, width=Cm(width_cm))  # 只给宽度 → 等比缩放
        if caption:
            self.add_figure_caption(caption, keep_with_next=bool(note),
                                    caption_en=caption_en)
        if note:
            self.add_note(note, align=WD_ALIGN_PARAGRAPH.CENTER)
        return p

    # ---------- 公式
    def add_equation(self, latex, number=None, show_number=True):
        self.eq_no += 1
        total = self.spec["table_width_cm"]
        widths = [2.0, total - 4.0, 2.0]
        table = self.doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_layout_fixed(table, total)
        remove_table_borders(table)
        set_col_widths(table, widths)
        forbid_row_split(table)
        cells = table.rows[0].cells
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        for cell in (cells[0], cells[2]):
            p = cell.paragraphs[0]
            para_format(p, first_line_indent=0, first_line_chars=0,
                        space_before=3, space_after=3, line_spacing=1.0)

        p = cells[1].paragraphs[0]
        para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                    first_line_chars=0, space_before=3, space_after=3,
                    line_spacing=1.0)
        add_latex_to_paragraph(p, latex)

        if show_number:
            num = number or "（%d）" % self.eq_no
            p = cells[2].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            self._run(p, num, size=self._s("body"), color=self._c("ink"))
        return table

    # ---------- 参考文献
    def add_references(self, items, heading="参考文献"):
        self.add_heading1(heading)
        for i, item in enumerate(items, 1):
            text = str(item).strip()
            if not re.match(r"^\[\d+\]", text):
                text = "[%d] %s" % (i, text)
            p = self.doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.2,
                        space_before=0, space_after=2,
                        left_indent=Cm(0.74), hanging_cm=0.74)
            self._run(p, text, size=self._s("reference"), color=self._c("ink"))

    # ---------- 摘要
    def add_abstract_zh(self, heading, paragraphs, keywords, page_break=False):
        self.add_heading1(heading, page_break=page_break, front=True)
        for text in paragraphs:
            self.add_body(text)
        if keywords:
            p = self.doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                        space_before=6, space_after=6,
                        first_line_indent=Pt(0), first_line_chars=0)
            self._run(p, "关键词：", self._f("head_cn"), self._f("head_en"),
                      self._s("keywords"), True, self._c("ink"))
            self._run(p, "；".join(keywords), self._f("body_cn"),
                      self._f("body_en"), self._s("keywords"), False,
                      self._c("ink"))

    def add_abstract_en(self, heading, paragraphs, keywords, page_break=False):
        if page_break:
            self.add_heading1(heading, page_break=True, front=True)
        else:
            self.add_heading2(heading)
        for text in paragraphs:
            p = self.doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                        space_before=0, space_after=6,
                        first_line_indent=Pt(0), first_line_chars=0)
            self._run(p, text, self._f("body_cn"), self._f("body_en"),
                      self._s("abstract_en"), color=self._c("ink"))
        if keywords:
            p = self.doc.add_paragraph()
            para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                        space_before=6, space_after=6,
                        first_line_indent=Pt(0), first_line_chars=0)
            self._run(p, "Keywords: ", self._f("head_cn"), self._f("head_en"),
                      self._s("keywords"), True, self._c("ink"))
            self._run(p, "; ".join(keywords), self._f("body_cn"),
                      self._f("body_en"), self._s("keywords"), False,
                      self._c("ink"))

    # ---------- 目录
    def add_toc(self, title="目　录", levels="1-1"):
        p = self.doc.add_paragraph()
        para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                    first_line_chars=0, space_before=0, space_after=18,
                    line_spacing=1.2, page_break_before=True)
        self._run(p, title, self._f("head_cn"), self._f("head_en"),
                  self._s("toc_title"), True, self._c("navy"))

        toc_p = self.doc.add_paragraph()
        para_format(toc_p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=0,
                    first_line_chars=0, line_spacing=1.5)
        font = (self._f("body_cn"), self._f("body_en"), 10.5, False,
                rgb(self._c("gray")))
        _field_part(toc_p, "begin", font=font)
        _field_part(toc_p, "instr", instr=' TOC \\o "%s" \\h \\z \\u ' % levels,
                    font=font)
        _field_part(toc_p, "separate", font=font)
        _field_part(toc_p, "text",
                    text="（目录域已插入：在 Word 中 Ctrl+A → F9，或右键“更新域”即可生成目录）",
                    font=font)
        _field_part(toc_p, "end", font=font)
        return toc_p

    # ---------- 页眉页脚
    def add_header_footer(self, meta):
        sec = self.doc.sections[0]
        title = meta.get("title", "")
        doc_type = meta.get("header_type") or meta.get("doc_type", "")
        header_text = meta.get("header_text") or (" · ".join(
            x for x in [title, doc_type] if x))

        hdr_p = sec.header.paragraphs[0]
        hdr_p.text = ""
        para_format(hdr_p, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line_indent=0,
                    first_line_chars=0, space_before=0, space_after=0,
                    line_spacing=1.0)
        _clear_tabs(hdr_p)
        self._run(hdr_p, header_text, size=self._s("header"),
                  color=self._c("gray"))

        ftr_p = sec.footer.paragraphs[0]
        ftr_p.text = ""
        para_format(ftr_p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
                    first_line_chars=0, space_before=0, space_after=0,
                    line_spacing=1.0)
        _clear_tabs(ftr_p)
        size, color = self._s("footer"), rgb(self._c("gray"))
        self._run(ftr_p, "第 ", size=size, color=self._c("gray"))
        add_field(ftr_p, "PAGE", placeholder="1", size=size, color=color)
        self._run(ftr_p, " 页 / 共 ", size=size, color=self._c("gray"))
        add_field(ftr_p, "NUMPAGES", placeholder="1", size=size, color=color)
        self._run(ftr_p, " 页", size=size, color=self._c("gray"))

    # ---------- 内容模型渲染
    def render(self, item):
        if not isinstance(item, dict):
            self.add_body(str(item))
            return
        level = item.get("level")
        typ = item.get("type")
        if typ in ("h1", "heading1", None) and level == 1:
            typ = "h1"
        elif typ in ("h2", "heading2") or level == 2:
            typ = "h2"

        if typ == "h1":
            self.add_heading1(item["text"], page_break=item.get("page_break", False),
                              front=item.get("front", False))
        elif typ == "h2":
            self.add_heading2(item["text"])
        elif typ == "p" or typ == "paragraph":
            texts = item["text"]
            if isinstance(texts, str):
                texts = [texts]
            for t in texts:
                self.add_body(t, indent_chars=item.get("indent_chars", 2))
        elif typ == "list":
            self.add_list(item["items"], start=item.get("start", 1))
        elif typ == "table":
            self.add_table(item.get("headers"), item.get("rows"),
                           caption=item.get("caption"),
                           caption_en=item.get("caption_en"),
                           note=item.get("note"),
                           widths=item.get("widths"),
                           aligns=_aligns(item.get("aligns")),
                           table_size=item.get("font_size"),
                           columns=item.get("columns"),
                           digits=item.get("digits"))
        elif typ == "figure":
            self.add_figure(item["path"], caption=item.get("caption"),
                            caption_en=item.get("caption_en"),
                            note=item.get("note"),
                            width_cm=item.get("width_cm"))
        elif typ == "equation":
            self.add_equation(item.get("latex", item.get("text", "")),
                              number=item.get("number"),
                              show_number=item.get("numbered", True))
        elif typ == "callout":
            self.add_callout(item.get("text", ""), item.get("title"))
        elif typ == "pagebreak":
            self.page_break()
        else:
            raise ValueError("未知的内容块类型：%r" % (item,))

    def build(self, model, out_path):
        self.setup()
        meta = model.get("meta", {})
        self.add_cover(meta)

        zh = model.get("abstract")
        en = model.get("abstract_en")
        if zh:
            self.add_abstract_zh(zh.get("heading", "摘　要"),
                                 _as_list(zh.get("paragraphs", [])),
                                 _as_list(zh.get("keywords", [])),
                                 page_break=True)
        if en:
            self.add_abstract_en(en.get("heading", "Abstract"),
                                 _as_list(en.get("paragraphs", [])),
                                 _as_list(en.get("keywords", [])),
                                 page_break=not zh)

        if model.get("toc", True):
            self.add_toc(model.get("toc_title", "目　录"),
                         model.get("toc_levels", "1-1"))

        first_h1 = True
        for item in model.get("sections", []):
            if isinstance(item, dict) and item.get("type") in ("h1", None) \
                    and item.get("level", 1) == 1:
                if first_h1 and model.get("body_page_break", True):
                    item = dict(item)
                    item.setdefault("page_break", True)
                    first_h1 = False
            self.render(item)

        refs = model.get("references") or []
        if refs:
            self.add_references(refs, model.get("references_heading", "参考文献"))

        self.add_header_footer(meta)
        enable_update_fields(self.doc)
        set_print_view(self.doc)

        out_path = os.path.abspath(out_path)
        os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
        self.doc.save(out_path)
        return out_path


def _as_list(value):
    if value is None:
        return []
    if isinstance(value, str):
        return [value]
    return list(value)


_ALIGN_MAP = {
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _aligns(values):
    if not values:
        return None
    return [_ALIGN_MAP.get(str(v).lower(), WD_ALIGN_PARAGRAPH.LEFT) for v in values]


_SUP = re.compile(r"\^\{([^}]*)\}")


def add_rich_text(paragraph, text, cn="宋体", en="Times New Roman", size=10.5,
                  bold=False, color=None):
    """写入一段文本，`^{a}` 渲染为上标（用于显著性标记与表注脚标）。"""
    pos = 0
    for m in _SUP.finditer(text):
        if m.start() > pos:
            set_run_font(paragraph.add_run(text[pos:m.start()]), cn, en, size,
                         bold, color)
        run = set_run_font(paragraph.add_run(m.group(1)), cn, en, size, bold, color)
        run.font.superscript = True
        pos = m.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), cn, en, size, bold, color)


def format_number(value, digits=None):
    """按列精度格式化。

    GB/T 7713.2—2022 5.5.4：表身中**空白**表示无此项或未测量，**—** 表示测量过
    而未发现，**0** 表示实测结果为零。因此 None/"" → 空白，不自动写"—"。
    """
    if value is None:
        return ""      # 5.5.4：空白 = 无此项或未测量；需要“—”时请在数据里显式写 "—"
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            return ""
        if digits is None:
            return text
        try:
            return format_number(float(text), digits)
        except ValueError:
            return text
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if digits is None:
            return ("%g" % value) if isinstance(value, float) else str(value)
        return ("%." + str(int(digits)) + "f") % float(value)
    return str(value)


def build_headers(columns):
    """从列元信息生成表头：'量名称/单位'（规范 3.3）。"""
    headers = []
    for col in columns:
        label = str(col.get("label", ""))
        unit = str(col.get("unit", "")).strip()
        headers.append("%s/%s" % (label, unit) if unit else label)
    return headers


def _clear_tabs(paragraph):
    """清掉 Header/Footer 样式自带的制表位，保证对齐可控。"""
    pPr = paragraph._p.get_or_add_pPr()
    remove_children(pPr, "tabs")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "clear")
    tab.set(qn("w:pos"), "4680")
    tabs.append(tab)
    tab2 = OxmlElement("w:tab")
    tab2.set(qn("w:val"), "clear")
    tab2.set(qn("w:pos"), "9360")
    tabs.append(tab2)
    insert_ordered(pPr, tabs, PPR_ORDER)


def build_from_model(model, out_path, spec_overrides=None, base_dir=None):
    builder = PaperBuilder(spec_overrides, base_dir=base_dir)
    out_path = builder.build(model, out_path)
    return builder, out_path
