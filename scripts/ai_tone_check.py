#!/usr/bin/env python
"""学术文风审计：按 academic-humanizer 的逻辑查 AI 痕迹与「主张-证据」不匹配。

逻辑来源（MIT 许可）：https://github.com/AIScientists-Dev/academic-humanizer
—— Layer 1 通用 AI 痕迹、Layer 2 学术 AI 痕迹（2.1–2.11）、Layer 3 保留清单、
Layer 4 主张-证据纪律、Layer 5 语体匹配、Layer 6 基金本子模式。
本文件只做**可枚举部分的审计**，不代替人工判断，也绝不改稿。

四步流程（与上游一致）：
  Read（读稿、认文体）→ Audit（列出命中位置与改法，不改）→
  Rewrite（由人或 agent 改）→ Report（改动清单）

用法：

    python scripts/ai_tone_check.py --docx out/report.docx
    python scripts/ai_tone_check.py --json work/content.json --audit
    python scripts/ai_tone_check.py --text draft.md --max-score 40

输出：按图层分组的问题清单（类型 + 段落号 + 原句 + 改法）+ 两个分数：
  tells  —— AI 痕迹密度（Layer 1 + 2）
  claims —— 主张强于证据的程度（Layer 4）
两者都低才算合格；分数只用于定位，不判断作者身份。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import statistics
import sys
from dataclasses import dataclass, field

import _console

CJK = r"\u4e00-\u9fa5"
_NUM = r"\d+(?:\.\d+)?"


# ---------------------------------------------------------------- 规则

@dataclass
class Rule:
    id: str
    layer: str          # "T1" 通用痕迹 / "T2" 学术痕迹 / "C" 主张-证据
    title: str
    pattern: str
    advice: str
    weight: int
    # True：句中含数字/百分比即视为"有证据"，不报（Layer 3 保留）
    needs_no_number: bool = False
    # True：只在该段最后一句命中才算（收尾套话）
    paragraph_final: bool = False


TELL_RULES: list[Rule] = [
    Rule(
        "T1.1", "T1", "意义拔高",
        r"标志(?:着|了)?(?:一个)?(?:重要|全新|关键)?的?(?:转折|里程碑|新阶段)|"
        r"具有(?:重要|深远|重大)?的?(?:理论|现实|实践)?(?:意义|价值|作用)|"
        r"奠定了?坚实的基础|划时代",
        "删掉评价，写清这个方法约束了什么、改变了哪个结论",
        4,
    ),
    Rule(
        "T1.2", "T1", "假深度的尾巴",
        r"[，,]\s*(?:从而|进而)?(?:凸显|彰显|体现|展现出?|揭示了?)(?:了)?(?:其|该|这一|这种)?"
        r"(?:重要|巨大|独特|深远)?的?(?:价值|潜力|优势|作用|意义)",
        "句尾补一句“这说明了什么”通常没信息；删掉或换成具体结论",
        3,
    ),
    Rule(
        "T1.3", "T1", "宣传化用词",
        r"赋能|助力|深耕|沉淀|护航|闭环|抓手|护城河|颠覆|革命性|璀璨|堪称",
        "换中性词：赋能→支持；助力→帮助；深耕→长期研究；闭环→反馈机制",
        3,
    ),
    Rule(
        "T1.4", "T1", "无来源的权威",
        r"(?:专家|学者|业界|业内|研究表明|有研究指出)\S{0,6}(?:认为|指出|表示|发现)",
        "给引用编号；给不出就删掉这句权威背书",
        4,
    ),
    Rule(
        "T1.5", "T1", "AI 高频空词",
        r"深入(?:探讨|剖析|挖掘)|多维(?:度)?|全链路|系统性(?:地)?(?:解决|提升)|"
        r"深度赋能|全方位|极大地|强有力地",
        "换成具体动作与对象：深入探讨→分析 X 的哪一步",
        2,
    ),
    Rule(
        "T1.6", "T1", "回避系动词",
        r"作为(?:一个)?(?:重要|关键|核心)的?(?:组成|环节|部分|载体)|充当着|"
        r"扮演着?(?:重要)?角色",
        "直接写“是/属于”：X 是 Y 的一部分",
        2,
    ),
    Rule(
        "T1.7", "T1", "否定式排比",
        r"不仅(?:仅)?是[^。；！？]{2,30}(?:更是|而且是|也是)|不是[^。；！？]{2,20}而是",
        "除非真在纠正一个误解，否则直接说它是什么",
        2,
    ),
    Rule(
        "T1.8", "T1", "三项式凑数",
        r"(?:(?:快速|准确|稳定|高效|可靠|灵活)(?:、|，)){2}"
        r"(?:快速|准确|稳定|高效|可靠|灵活)",
        "三个并列形容词通常可删两个，或换成可测指标",
        2,
    ),
    Rule(
        "T1.9", "T1", "同义反复换词",
        r"换言之|换句话说|也就是说[^。；]{0,6}(?:重复|同样)|即所谓的",
        "同一概念全文用一个词，不要换着说",
        1,
    ),
    Rule(
        "T1.10", "T1", "填充套话",
        r"值得(?:一提|注意)的是|众所周知|不言而喻|毋庸置疑",
        "直接删；重要的话让句子本身显示出来",
        3,
    ),
    Rule(
        "T1.11", "T1", "长句堆砌",
        r"(?:[^，,。；！？]{10,}[，,]){3,}",
        "一句话一个意思：3 个以上长分句就拆句，删掉不承重的从句",
        2,
    ),
    Rule(
        "T1.12", "T1", "破折号",
        r"——",
        "学术正文不用破折号，改用逗号、冒号、括号或另起一句",
        2,
    ),
]

ACADEMIC_RULES: list[Rule] = [
    Rule(
        "T2.1", "T2", "过度断言动词",
        r"证明|证实|确保了|保证了|完全(?:消除|解决|避免)|彻底(?:解决|改变)",
        "实证工作只能说“显示/提供了证据”；动词强度不得超过证据",
        4,
        needs_no_number=True,
    ),
    Rule(
        "T2.2", "T2", "显著性炒作",
        r"开辟了?(?:新的)?(?:道路|方向|范式)|为[^。；]{0,10}铺平了?道路|"
        r"有望(?:彻底)?(?:改变|颠覆|革新)|填补了?(?:空白|空白点)|"
        r"(?:显著|大幅|明显|极大)地?(?:提升|提高|改善|增强|降低|减少|缩短)",
        "写清改进了哪个失效模式、提升了多少；补数字与检验",
        4,
        needs_no_number=True,
    ),
    Rule(
        "T2.3", "T2", "空泛强化词",
        r"广泛的?(?:实验|测试|评估|数据)|全面的?(?:实验|评估|分析)|"
        r"大量(?:的)?(?:实验|数据)|多种(?:数据|场景|方法)",
        "举出具体数据集、样本量、场景数，而不是“广泛”“大量”",
        3,
        needs_no_number=True,
    ),
    Rule(
        "T2.4", "T2", "新颖性填充",
        r"首次(?:提出|实现|研究)|据(?:我们|笔者)所知|据(?:我们|笔者)了解|"
        r"填补了?(?:该|这一)?空白",
        "改为“先前工作在 X 上未覆盖 Y”；新颖性靠对比，不靠宣言",
        3,
    ),
    Rule(
        "T2.5", "T2", "模板化开头",
        r"随着[^。；]{0,20}的(?:快速|迅速|不断)?(?:发展|进步|普及)[^。；]{0,20}[，,]|"
        r"近年来[，,][^。；]{0,20}(?:受到|引起)[^。；]{0,10}(?:广泛)?关注",
        "删掉铺垫，直接从具体问题或失效模式开头",
        3,
    ),
    Rule(
        "T2.6", "T2", "连接词连用",
        r"^(?:此外|同时|另外|而且|并且|然后|其次)[^。；]{0,10}[，,]",
        "连续几句都用连接词开头会像清单；让逻辑自己承接",
        2,
    ),
    Rule(
        "T2.7", "T2", "贡献清单套话",
        r"本文(?:的)?(?:主要)?贡献(?:如下|包括)[：:]?\s*(?:首先|第一|\(1\)|（1）)",
        "每条贡献写成一个具体结果（指标+对比+条件），不要复述摘要",
        3,
    ),
    Rule(
        "T2.8", "T2", "文献堆砌",
        r"许多(?:研究|方法|工作)(?:已经)?(?:提出|存在|表明)|大量文献",
        "点出最近的一两篇并说明差异，而不是一串引用",
        3,
    ),
    Rule(
        "T2.9", "T2", "模糊限定",
        r"(?:较为|一定程度上|某种程度上|相对来说|比较)(?:好|差|高|低|明显|稳定|准确)?",
        "量化或删掉：给数值、区间与条件",
        2,
        needs_no_number=True,
    ),
    Rule(
        "T2.10", "T2", "强调套话",
        r"(?:值(?:得)?(?:特别)?(?:指出|注意)|应当(?:强调|指出)|尤其(?:值得)?注意)[的是]?[，,]?",
        "删掉；若真重要，用数据和结论本身说明",
        3,
    ),
    Rule(
        "T2.11", "T2", "从句堆叠",
        r"(?:[^。；！？]{20,})(?:，|,)?(?:其|该|这一)(?:[^。；！？]{10,})"
        r"(?:，|,)(?:这|该)(?:[^。；！？]{10,})",
        "拆句；一个句子只承担一个判断",
        2,
    ),
]

CLAIM_RULES: list[Rule] = [
    Rule(
        "C1", "C", "主张无证据指针",
        r"(?:本文|该方法|本方法|该模型)[^。；！？]{0,30}"
        r"(?:优于|高于|超过|领先|更好|更优|更准|更快)",
        "实证主张要落到数字、图表或文献；补“[图3]”“[表2]”“[12]”，或降调",
        4,
        needs_no_number=True,
    ),
    Rule(
        "C2", "C", "量级模糊",
        r"(?:提升|提高|改善|降低|减少|缩短|增长)(?:了)?(?:不少|很多|许多|一些|一倍多)",
        "给区间与统计口径，如“提升 3—6 个百分点（n=5，p<0.01）”",
        3,
    ),
    Rule(
        "C3", "C", "比较对象缺失",
        r"优于(?:现有|已有|传统|主流)(?:方法|方案)",
        "指明是哪一个基线、在哪一指标上、差多少",
        3,
    ),
    Rule(
        "C4", "C", "结论强于设计",
        r"(?:可|能够)(?:广泛)?(?:应用|推广|适用)于[^。；！？]{0,12}(?:所有|各种|各类|任何)",
        "把适用范围写成实测条件；“所有”类表述通常没有证据支撑",
        3,
    ),
]

# —— Layer 6：基金本子（结构 + 弱招），靠结构判断，不是逐句正则
PROPOSAL_MARKERS = {
    "nsf": {
        "label": "NSF",
        "structure": [("Project Summary", r"project\s*summary|项目摘要"),
                      ("Overview", r"overview|概述"),
                      ("Intellectual Merit", r"intellectual\s*merit|智力价值|学术价值"),
                      ("Broader Impacts", r"broader\s*impacts|广泛影响|社会影响")],
    },
    "nih": {
        "label": "NIH",
        "structure": [("Specific Aims", r"specific\s*aims|具体目标"),
                      ("Significance", r"significance|意义"),
                      ("Innovation", r"innovation|创新"),
                      ("Approach", r"approach|研究方案")],
    },
}
AIM_RE = re.compile(r"(?:Aim|Thrust|目标)\s*[0-9一二三四五六]?\s*[：:]")
METHOD_AIM_RE = re.compile(
    r"(?:Aim|Thrust|目标)\s*[0-9一二三四五六]?\s*[：:]\s*"
    r"(?:应用|采用|使用|引入|构建|开发|实现|搭建)\S{0,16}(?:方法|模型|算法|框架|系统|平台)")
DOMINO_RE = re.compile(
    r"(?:若|如果|一旦|如)\s*(?:Aim|目标)\s*[0-9一二三]"
    r"(?:\s*(?:与|和|及|、)\s*(?:Aim|目标)?\s*[0-9一二三])?\s*(?:成功|顺利|完成)"
    r"|(?:Aim|目标)\s*2\s*(?:依赖|基于)\s*(?:Aim|目标)\s*1")
HEDGED_HYPOTHESIS_RE = re.compile(
    r"(?:central\s*hypothesis|中心假设)[^。；！？]{0,40}"
    r"(?:explore|whether|可能|或许|探索|是否)")
FEASIBILITY_RE = re.compile(
    r"前期|初步|已有|已获|已发表|预实验|pilot|preliminary|prior\s+work|"
    r"Fig\.?\s*\d|图\s*\d|letter|支持信|合作|团队")
AMBITION_RE = re.compile(r"首次|突破|颠覆|填补空白|有望|变革性|transformative|pioneer")
BOILERPLATE_BI_RE = re.compile(r"(?:传播|宣讲|推广|科普|培养学生|开设课程)")


def _paras_text(paragraphs: list[tuple[str, bool]]) -> tuple[str, list[str]]:
    rows = [t for t, h in paragraphs if not h]
    return "\n".join(rows), rows


def analyze_proposal(paragraphs: list[tuple[str, bool]],
                     kind: str | None = None) -> dict:
    """Layer 6：查结构完整性、aims 质量与「主张 ↔ 可行性」。

    kind: "nsf" / "nih"；None 时按文中出现的标记自动判断。
    """
    text, rows = _paras_text(paragraphs)
    if kind is None:
        kind = "nih" if re.search(r"specific\s*aims|具体目标", text, re.I) else "nsf"
    spec = PROPOSAL_MARKERS[kind]

    issues: list[tuple[str, str, str]] = []
    missing = [name for name, pat in spec["structure"]
               if not re.search(pat, text, re.I)]
    # Aims 页只承载 Aims；Significance/Innovation/Approach 属正文，不作缺失项
    aims_page = bool(re.search(r"specific\s*aims|具体目标", text, re.I))
    if kind == "nih" and aims_page:
        missing = [m for m in missing
                   if m not in ("Significance", "Innovation", "Approach")]
    if missing:
        issues.append(("P1", "结构不完整",
                       "%s 缺少：%s" % (spec["label"], "、".join(missing))))

    aims = AIM_RE.findall(text)
    n_aims = len(aims) if isinstance(aims, list) else 0
    if aims and not (2 <= n_aims <= 3):
        issues.append(("P2", "aims 数量异常",
                       "检测到 %d 个 aim，NSF/NIH 通常 2—3 个" % n_aims))

    method_aims = METHOD_AIM_RE.findall(text)
    if method_aims:
        issues.append(("P3", "方法当目标",
                       "%d 处 aim 以方法动词开头（应用/采用/构建…），应改为问题或结果导向"
                       % len(method_aims)))

    if DOMINO_RE.search(text):
        issues.append(("P4", "aims 有多米诺依赖",
                       "Aim 之间出现“若 Aim 1 成功”式依赖；应写成平行、各自可交付"))

    if HEDGED_HYPOTHESIS_RE.search(text):
        issues.append(("P6", "中心假设含糊",
                       "中心假设处出现“探索/是否/可能”；中心假设应可证伪"))

    # P5：雄心句附近 40 字内无可行性证据
    weak = []
    for m in AMBITION_RE.finditer(text):
        window = text[max(0, m.start() - 40): m.end() + 40]
        if not FEASIBILITY_RE.search(window):
            weak.append(text[max(0, m.start() - 12): m.end() + 12].replace("\n", " "))
    if weak:
        issues.append(("P5", "雄心缺可行性支撑",
                       "%d 处大胆主张周边无前期数据/前期工作/合作方证据：%s"
                       % (len(weak), "；".join(weak[:2]))))

    if BOILERPLATE_BI_RE.search(text) and not re.search(r"\d", text):
        issues.append(("P7", "Broader Impacts 模板化",
                       "出现传播/宣讲/培养学生，但全文无具体数字或名称"))

    return {"kind": kind, "label": spec["label"], "issues": issues,
            "aims": len(aims) if isinstance(aims, list) else 0}


ALL_RULES = TELL_RULES + ACADEMIC_RULES + CLAIM_RULES
RULE_BY_ID = {r.id: r for r in ALL_RULES}

# Layer 3 保留清单：证据型模糊限制属于正确写法，不报
PRESERVE_HINTS = re.compile(
    r"(?:提示|表明|说明|与[^。；]{0,6}一致|推测|假设|可能|似乎|有待)"
    r"[^。；！？]{0,24}(?:\d|图|表|\[)")


# ---------------------------------------------------------------- 文本处理

CONNECTIVES = r"因此|然而|此外|同时|首先|其次|最后|综上|总之|另外|需要指出|由此可见"
CITATION = re.compile(r"\[\s*\d+(?:\s*[-,，]\s*\d+)*\s*\]")
QUOTED = re.compile(r"「[^」]{0,80}」|“[^”]{0,80}”|\"[^\"]{0,80}\"|‘[^’]{0,80}’|`[^`]{0,80}`")
SENT_SPLIT = re.compile(r"[。！？；!?;]+")
CAPTION = re.compile(r"^(?:图|表)\s*\d+|^注：")
REF_HEADING = re.compile(r"^(?:参考文献|References)$", re.I)
HEADING_STYLE = re.compile(r"^(?:Heading|标题)\s*\d", re.I)


@dataclass
class Hit:
    rule_id: str
    layer: str
    title: str
    para_no: int
    excerpt: str
    advice: str


@dataclass
class Report:
    source: str
    chars: int = 0
    paragraphs: int = 0
    sentences: int = 0
    hits: list = field(default_factory=list)
    metrics: dict = field(default_factory=dict)

    def _score(self, layers: set[str]) -> int:
        if not self.chars:
            return 0
        weight = 0.0
        counted: dict[str, int] = {}
        for hit in self.hits:
            if hit.layer not in layers:
                continue
            rule = RULE_BY_ID[hit.rule_id]
            n = counted.get(hit.rule_id, 0)
            counted[hit.rule_id] = n + 1
            # 同一规则第 3 次起降权，避免一个坏习惯把分数拉爆
            weight += rule.weight if n < 2 else rule.weight * 0.5
        denominator = max(self.chars, 500) / 1000.0
        return int(min(100, round(weight / denominator * 4)))

    @property
    def tells_score(self) -> int:
        """AI 痕迹密度（Layer 1 + 2）。"""
        return self._score({"T1", "T2"})

    @property
    def claims_score(self) -> int:
        """主张强于证据的程度（Layer 4）。"""
        return self._score({"C"})

    @property
    def score(self) -> int:
        """总分，供 --max-score 使用。"""
        return min(100, self.tells_score + self.claims_score)


def strip_noise(text: str) -> str:
    """剥离引用编号、引语、行内代码，避免误报。"""
    text = CITATION.sub("", text)
    text = QUOTED.sub("", text)
    return text.replace("`", "")


def split_sentences(text: str) -> list[str]:
    return [s.strip() for s in SENT_SPLIT.split(text) if s.strip()]


# ---------------------------------------------------------------- 采集

def collect_paragraphs_docx(path: str) -> tuple[list[tuple[str, bool]], str]:
    from docx import Document

    doc = Document(path)
    out: list[tuple[str, bool]] = []
    in_refs = False
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        if REF_HEADING.match(text):
            in_refs = True
            continue
        if in_refs or CAPTION.match(text):
            continue
        out.append((text, bool(HEADING_STYLE.match(style))))
    return out, os.path.basename(path)


def collect_paragraphs_model(path: str) -> tuple[list[tuple[str, bool]], str]:
    with open(os.path.abspath(path), encoding="utf-8-sig") as fh:
        model = json.load(fh)
    out: list[tuple[str, bool]] = []
    for item in model.get("sections", []):
        if not isinstance(item, dict):
            out.append((str(item), False))
            continue
        if item.get("level") in (1, 2, 3) or item.get("type") in ("h1", "h2", "h3"):
            out.append((str(item.get("text", "")), True))
        elif item.get("type") in ("p", "paragraph"):
            texts = item.get("text", "")
            if isinstance(texts, str):
                texts = [texts]
            for t in texts:
                out.append((str(t), False))
        elif item.get("type") == "callout":
            out.append((str(item.get("text", "")), False))
    for key in ("abstract", "abstract_en"):
        block = model.get(key) or {}
        for t in block.get("paragraphs", []) or []:
            out.append((str(t), False))
    return out, os.path.basename(path)


def collect_paragraphs_text(path: str) -> tuple[list[tuple[str, bool]], str]:
    with open(path, encoding="utf-8") as fh:
        raw = fh.read()
    out = []
    for block in re.split(r"\n\s*\n", raw):
        block = block.strip()
        if not block:
            continue
        if block.startswith("#"):
            out.append((block.lstrip("# ").strip(), True))
        else:
            out.append((block.replace("\n", ""), False))
    return out, os.path.basename(path)


# ---------------------------------------------------------------- 检测

def analyze(paragraphs: list[tuple[str, bool]], source: str) -> Report:
    report = Report(source=source)
    report.paragraphs = len(paragraphs)

    all_chars = 0
    all_sentences: list[int] = []
    connective_count = 0
    number_count = 0

    for idx, (raw, is_heading) in enumerate(paragraphs, start=1):
        if is_heading:
            continue
        text = strip_noise(raw)
        all_chars += len(text)
        sentences = split_sentences(text)
        all_sentences.extend(len(s) for s in sentences)
        connective_count += len(re.findall(CONNECTIVES, text))
        number_count += len(re.findall(_NUM, text))

        for rule in ALL_RULES:
            for s_idx, sentence in enumerate(sentences):
                if rule.needs_no_number and re.search(_NUM, sentence):
                    continue          # 有数字=有证据（Layer 3 保留）
                if rule.layer == "C" and PRESERVE_HINTS.search(sentence):
                    continue          # 证据型模糊限制不报
                if rule.paragraph_final and s_idx != len(sentences) - 1:
                    continue
                if not re.search(rule.pattern, sentence):
                    continue
                report.hits.append(Hit(rule.id, rule.layer, rule.title, idx,
                                       _excerpt(sentence), rule.advice))
                break                 # 同一段同一规则只报一次

    report.chars = all_chars
    report.sentences = len(all_sentences)
    report.metrics = {
        "句长均值": round(statistics.mean(all_sentences), 1) if all_sentences else 0,
        "句长变异": round(_cv(all_sentences), 2),
        "连接词/千字": round(connective_count / all_chars * 1000, 1) if all_chars else 0,
        "数字/千字": round(number_count / all_chars * 1000, 1) if all_chars else 0,
    }
    return report


def _cv(values: list[int]) -> float:
    if len(values) < 2:
        return 0.0
    mean = statistics.mean(values)
    return statistics.pstdev(values) / mean if mean else 0.0


def _excerpt(text: str, limit: int = 46) -> str:
    text = text.strip()
    return text if len(text) <= limit else text[:limit] + "…"


# ---------------------------------------------------------------- 输出

LABELS = [(60, "重"), (40, "明显"), (20, "轻微"), (0, "干净")]


def grade(score: int) -> str:
    for threshold, label in LABELS:
        if score >= threshold:
            return label
    return "干净"


def render(report: Report) -> str:
    lines = ["学术文风审计：%s" % report.source,
             "分析 %d 字 / %d 段 / %d 句" % (report.chars, report.paragraphs,
                                             report.sentences),
             "-" * 72]
    tells, claims = report.tells_score, report.claims_score
    lines.append("AI 痕迹 %d / 100（%s）｜主张-证据 %d / 100（%s）"
                 % (tells, grade(tells), claims, grade(claims)))
    lines.append("两者都低才算合格。分数只用于定位，不判断作者身份。")
    m = report.metrics
    lines.append("句长均值 %s ｜ 句长变异 %s ｜ 连接词 %s/千字 ｜ 数字 %s/千字"
                 % (m.get("句长均值"), m.get("句长变异"),
                    m.get("连接词/千字"), m.get("数字/千字")))
    lines.append("-" * 72)

    if not report.hits:
        lines.append("没有命中可枚举的规则。仍建议按 Layer 4 逐条核对主张与证据：")
        lines.append("  references/anti-ai-academic-zh.md")
        return "\n".join(lines)

    grouped: dict[str, list[Hit]] = {}
    for hit in report.hits:
        grouped.setdefault(hit.rule_id, []).append(hit)

    for layer, label in (("T1", "Layer 1 通用 AI 痕迹"),
                         ("T2", "Layer 2 学术 AI 痕迹"),
                         ("C", "Layer 4 主张-证据")):
        rules = [r for r in ALL_RULES if r.layer == layer and r.id in grouped]
        if not rules:
            continue
        lines.append("【%s】" % label)
        for rule in rules:
            hits = grouped[rule.id]
            lines.append("  [%s] %s ×%d" % (rule.id, rule.title, len(hits)))
            for hit in hits[:6]:
                where = "全篇" if hit.para_no == 0 else "¶%d" % hit.para_no
                lines.append("     %-6s %s" % (where, hit.excerpt))
            if len(hits) > 6:
                lines.append("     … 另有 %d 处" % (len(hits) - 6))
            if rule.advice:
                lines.append("     改法：%s" % rule.advice)
        lines.append("")

    lines.append("下一步：按 Layer 3 逐句核对——有证据的限制、被动语态、“我们”、"
                 "术语与公式不动；数字、公式、引用一个都不许改。")
    return "\n".join(lines)


def render_proposal(result: dict) -> str:
    label = result.get("label", "?")
    lines = ["【Layer 6 基金本子（%s）】" % label,
             "  检测到 aim：%d 个" % result.get("aims", 0)]
    issues = result.get("issues") or []
    if not issues:
        lines.append("  OK   结构、aims 与可行性未见可自动判定的问题")
        lines.append("  仍需人工核对：中心假设是否可证伪、每个 aim 是否独立可交付、"
                     "团队记录是否真实（绝不编造资助/结果/伙伴/推荐信）")
        return "\n".join(lines)
    for rid, title, detail in issues:
        lines.append("  WARN [%s] %s" % (rid, title))
        lines.append("        %s" % detail)
    lines.append("  改法：按 references/anti-ai-academic-zh.md Layer 6.2—6.5 逐条补齐")
    return "\n".join(lines)


def render_audit(report: Report) -> str:
    """Audit 模式：一行一条，便于逐条处理。"""
    if not report.hits:
        return "无命中。"
    rows = ["层\t编号\t类型\t位置\t改法"]
    for hit in report.hits:
        rows.append("%s\t%s\t%s\t¶%d\t%s" % (hit.layer, hit.rule_id, hit.title,
                                            hit.para_no, hit.advice))
    return "\n".join(rows)


def main(argv=None):
    _console.setup()
    ap = argparse.ArgumentParser(
        description="按 academic-humanizer 的逻辑审计学术文风（只报警，不改稿）")
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--docx", help="待检测的 docx")
    src.add_argument("--json", help="内容模型 JSON")
    src.add_argument("--text", help="纯文本 / Markdown")
    ap.add_argument("--max-score", type=int, default=None,
                    help="总分阈值（超出退出码 1）")
    ap.add_argument("--json-out", action="store_true", help="以 JSON 输出")
    ap.add_argument("--audit", action="store_true",
                    help="只输出审计清单（层/编号/类型/位置/改法）")
    ap.add_argument("--mode", choices=["paper", "proposal"], default="paper",
                    help="paper=论文（默认）；proposal=基金本子（追加 Layer 6 检查）")
    ap.add_argument("--agency", choices=["nsf", "nih"], default=None,
                    help="proposal 模式下指定机构；省略则自动判断")
    args = ap.parse_args(argv)

    if args.docx:
        paragraphs, name = collect_paragraphs_docx(args.docx)
    elif args.json:
        paragraphs, name = collect_paragraphs_model(args.json)
    else:
        paragraphs, name = collect_paragraphs_text(args.text)

    report = analyze(paragraphs, name)
    proposal = (analyze_proposal(paragraphs, args.agency)
                if args.mode == "proposal" else None)
    if not report.chars and not (proposal and proposal["issues"]):
        print("没有可分析的正文段落（检查输入文件）", file=sys.stderr)
        return 2

    if args.json_out:
        print(json.dumps({
            "source": report.source,
            "tells_score": report.tells_score,
            "claims_score": report.claims_score,
            "score": report.score,
            "grade": grade(report.score),
            "chars": report.chars,
            "sentences": report.sentences,
            "metrics": report.metrics,
            "mode": args.mode,
            "proposal": proposal,
            "hits": [h.__dict__ for h in report.hits],
        }, ensure_ascii=False, indent=2))
    elif args.audit:
        print(render_audit(report))
        if proposal:
            print()
            print(render_proposal(proposal))
    else:
        print(render(report))
        if proposal:
            print()
            print(render_proposal(proposal))

    if args.max_score is not None and report.score > args.max_score:
        print("\n超过阈值 %d（实际 %d）。" % (args.max_score, report.score),
              file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
