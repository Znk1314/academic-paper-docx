"""set_view.py 测试：只改 settings.xml，其它部件必须一字不动。"""

from __future__ import annotations

import zipfile

import docx_builder as db
import set_view
from conftest import w
from docx import Document


def _parts(path):
    with zipfile.ZipFile(path) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def test_view_rewritten(built):
    path = built[1]
    doc = Document(path)
    db._set_setting(doc, "view", {"val": "web"})
    db._set_setting(doc, "zoom", {"val": "100"})
    doc.save(path)
    assert set_view.read_view(path) == ("web", None)

    set_view.set_print_view(path)
    assert set_view.read_view(path) == ("print", "100")


def test_only_settings_xml_changes(built):
    path = built[1]
    # 模拟 Word 保存后把视图改回 Web 版式
    doc = Document(path)
    db._set_setting(doc, "view", {"val": "web"})
    doc.save(path)

    before = _parts(path)
    set_view.set_print_view(path)
    after = _parts(path)

    assert before.keys() == after.keys()
    changed = {name for name in before if before[name] != after[name]}
    assert changed == {"word/settings.xml"}


def test_check_mode_exit_codes(built):
    path = built[1]
    set_view.set_print_view(path)
    assert set_view.main([path, "--check"]) == 0

    doc = Document(path)
    db._set_setting(doc, "view", {"val": "web"})
    doc.save(path)
    assert set_view.main([path, "--check"]) == 1


def test_zoom_value_set(built):
    path = built[1]
    set_view.set_print_view(path, zoom=90)
    with zipfile.ZipFile(path) as zf:
        from lxml import etree
        root = etree.fromstring(zf.read("word/settings.xml"))
    assert root.find(w("zoom")).get(w("percent")) == "90"


def test_invalid_docx_raises(tmp_path):
    bogus = tmp_path / "not-a-docx.docx"
    bogus.write_bytes(b"PK\x03\x04junk")
    try:
        set_view.set_print_view(str(bogus))
    except Exception:
        return
    raise AssertionError("对非法文件应当报错")
