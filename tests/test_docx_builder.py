"""格式库测试：把规范的每条硬指标变成断言。"""

from __future__ import annotations

import os

import docx_builder as db
import pytest
from conftest import DEMO_FIGURE, demo_model, localname, m, w
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

# --------------------------------------------------------------- 页面 / 视图

def test_page_setup(built_doc):
    sec = built_doc.sections[0]
    # docx 内部以 twips 存页面尺寸，换算回 EMU 会有 <0.01 mm 的舍入
    assert sec.page_width.cm == pytest.approx(21.0, abs=0.01)
    assert sec.page_height.cm == pytest.approx(29.7, abs=0.01)
    assert sec.top_margin.cm == pytest.approx(2.4, abs=0.01)
    assert sec.bottom_margin.cm == pytest.approx(2.2, abs=0.01)
    assert sec.left_margin.cm == pytest.approx(2.6, abs=0.01)
    assert sec.right_margin.cm == pytest.approx(2.4, abs=0.01)
    assert sec.header_distance.cm == pytest.approx(1.2, abs=0.01)
    assert sec.footer_distance.cm == pytest.approx(1.0, abs=0.01)


def test_first_page_header_footer_is_different(built_doc):
    sectPr = built_doc.sections[0]._sectPr
    assert sectPr.find(w("titlePg")) is not None


def test_print_view_is_default(built_doc):
    settings = built_doc.settings.element
    assert settings.find(w("view")).get(w("val")) == "print"
    assert settings.find(w("zoom")).get(w("percent")) == "100"
    assert settings.find(w("updateFields")).get(w("val")) == "true"


# --------------------------------------------------------------- 样式

def test_normal_style(built_doc):
    st = built_doc.styles["Normal"]
    rPr = st.element.find(w("rPr"))
    rFonts = rPr.find(w("rFonts"))
    assert rFonts.get(w("eastAsia")) == "宋体"
    assert rFonts.get(w("ascii")) == "Times New Roman"
    assert rPr.find(w("sz")).get(w("val")) == "24"
    pPr = st.element.find(w("pPr"))
    assert pPr.find(w("spacing")).get(w("line")) == "360"
    assert pPr.find(w("ind")).get(w("firstLine")) == "480"
    assert pPr.find(w("jc")).get(w("val")) == "both"


def test_normal_style_has_no_theme_fonts(built_doc):
    rFonts = built_doc.styles["Normal"].element.find(w("rPr")).find(w("rFonts"))
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        assert rFonts.get(w(attr)) is None


def test_heading1_style(built_doc):
    st = built_doc.styles["Heading 1"]
    rPr = st.element.find(w("rPr"))
    pPr = st.element.find(w("pPr"))
    assert rPr.find(w("sz")).get(w("val")) == "32"
    assert rPr.find(w("color")).get(w("val")) == "1F3A5F"
    assert rPr.find(w("b")) is not None
    bottom = pPr.find(w("pBdr")).find(w("bottom"))
    assert bottom.get(w("sz")) == "6"
    assert bottom.get(w("color")) == "0072B2"
    assert pPr.find(w("keepNext")) is not None


def test_heading2_style(built_doc):
    st = built_doc.styles["Heading 2"]
    rPr = st.element.find(w("rPr"))
    pPr = st.element.find(w("pPr"))
    assert rPr.find(w("sz")).get(w("val")) == "26"
    assert rPr.find(w("color")).get(w("val")) == "0072B2"
    assert pPr.find(w("keepNext")) is not None


def test_front_heading_is_not_in_toc(built_doc):
    """摘要用 Front Heading：外观同一级标题，但大纲级别 9 → 不进目录。"""
    style = next(s for s in built_doc.styles if s.name == "Front Heading")
    pPr = style.element.find(w("pPr"))
    assert pPr.find(w("outlineLvl")).get(w("val")) == "9"
    assert pPr.find(w("pBdr")) is not None


# --------------------------------------------------------------- 正文与列表

def test_body_paragraph_indent(built_doc):
    body = [p for p in built_doc.paragraphs
            if p.style.name == "Normal" and p.text.startswith("这是第一段")]
    assert body, "没有找到正文段落"
    ind = body[0]._p.find(w("pPr")).find(w("ind"))
    assert ind.get(w("firstLineChars")) == "200"


def test_list_items_use_fullwidth_numbering(built_doc):
    texts = [p.text for p in built_doc.paragraphs]
    assert "（1）要点一" in texts
    assert "（2）要点二" in texts
    for p in built_doc.paragraphs:
        assert "•" not in p.text
        pPr = p._p.find(w("pPr"))
        assert pPr is None or pPr.find(w("numPr")) is None


# --------------------------------------------------------------- 三线表

def _data_tables(doc):
    out = []
    for tbl in doc.element.body.iter(w("tbl")):
        borders = tbl.find(w("tblPr")).find(w("tblBorders"))
        vals = {localname(e): e.get(w("val")) for e in borders}
        if vals.get("top") == "single":
            out.append(tbl)
    return out


def test_three_line_table_borders(built_doc):
    tables = _data_tables(built_doc)
    assert len(tables) == 1
    tbl = tables[0]
    borders = {localname(e): e for e in tbl.find(w("tblPr")).find(w("tblBorders"))}
    assert borders["top"].get(w("sz")) == "12"
    assert borders["bottom"].get(w("sz")) == "12"
    for edge in ("left", "right", "insideH", "insideV"):
        assert borders[edge].get(w("val")) == "nil"


def test_table_header_is_bold_and_repeated(built_doc):
    tbl = _data_tables(built_doc)[0]
    header_tr = tbl.findall(w("tr"))[0]
    assert header_tr.find(w("trPr")).find(w("tblHeader")) is not None
    first_cell = header_tr.find(w("tc"))
    bottom = first_cell.find(w("tcPr")).find(w("tcBorders")).find(w("bottom"))
    assert bottom.get(w("sz")) == "6"
    run = first_cell.find(w("p")).find(w("r"))
    assert run.find(w("rPr")).find(w("b")) is not None


def test_table_rows_cannot_split(built_doc):
    tbl = _data_tables(built_doc)[0]
    for tr in tbl.findall(w("tr")):
        assert tr.find(w("trPr")).find(w("cantSplit")) is not None


def test_table_width_is_15_5_cm(built_doc):
    tbl = _data_tables(built_doc)[0]
    tblW = tbl.find(w("tblPr")).find(w("tblW"))
    assert tblW.get(w("type")) == "dxa"
    assert abs(int(tblW.get(w("w"))) - 8788) <= 8


def test_table_caption_above_table(built_doc):
    body = list(built_doc.element.body)
    tbl = _data_tables(built_doc)[0]
    idx = body.index(tbl)
    prev = body[idx - 1]
    text = "".join(t.text or "" for t in prev.iter(w("t")))
    assert text.startswith("表1")
    pPr = prev.find(w("pPr"))
    assert pPr.find(w("keepNext")) is not None
    assert pPr.find(w("jc")).get(w("val")) == "center"


def test_caption_numbering_is_renumbered(basic_model, tmp_path):
    """标题里写错编号也要被纠正为连续编号。"""
    basic_model["sections"] = [
        {"level": 1, "text": "1 引言"},
        {"type": "table", "caption": "表7 乱写的编号", "headers": ["A"], "rows": [["1"]]},
        {"type": "table", "caption": "另一个表", "headers": ["A"], "rows": [["1"]]},
    ]
    _, path = db.build_from_model(basic_model, str(tmp_path / "n.docx"))
    doc = Document(path)
    caps = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("表")]
    assert caps[0].startswith("表1")
    assert caps[1].startswith("表2")


# --------------------------------------------------------------- 图

def test_figure_is_proportional_and_kept_with_caption(basic_model, tmp_path):
    if not os.path.exists(DEMO_FIGURE):
        pytest.skip("缺少示例图片")
    basic_model["sections"] = [
        {"level": 1, "text": "1 引言"},
        {"type": "figure", "path": DEMO_FIGURE, "caption": "示例图", "width_cm": 14.0},
        {"type": "p", "text": "如图1所示。"},
    ]
    _, path = db.build_from_model(basic_model, str(tmp_path / "f.docx"))
    doc = Document(path)
    shape = doc.inline_shapes[0]
    assert abs(shape.width.cm - 14.0) < 0.05
    from PIL import Image
    with Image.open(DEMO_FIGURE) as img:
        native = img.width / img.height
    assert abs(shape.width / shape.height - native) / native < 0.01
    fig_para = next(p for p in doc.paragraphs
                    if p._p.findall(".//" + w("drawing")))
    assert fig_para._p.find(w("pPr")).find(w("keepNext")) is not None


def test_figure_too_wide_is_clamped_with_warning(basic_model, tmp_path):
    if not os.path.exists(DEMO_FIGURE):
        pytest.skip("缺少示例图片")
    basic_model["sections"] = [
        {"level": 1, "text": "1 引言"},
        {"type": "figure", "path": DEMO_FIGURE, "caption": "超宽图", "width_cm": 20.0},
    ]
    builder, path = db.build_from_model(basic_model, str(tmp_path / "w.docx"))
    assert builder.warnings
    shape = Document(path).inline_shapes[0]
    assert abs(shape.width.cm - 15.5) < 0.05


# --------------------------------------------------------------- 公式

def test_equation_table_layout(built_doc):
    eq_tables = [t for t in built_doc.tables if len(t.columns) == 3 and len(t.rows) == 1]
    assert eq_tables, "没有找到公式表格"
    cells = eq_tables[0].rows[0].cells
    assert cells[1]._tc.find(".//" + m("oMath")) is not None
    assert cells[2].text.strip() == "（1）"
    assert cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    borders = eq_tables[0]._tbl.find(w("tblPr")).find(w("tblBorders"))
    assert all(e.get(w("val")) == "nil" for e in borders)


def test_equation_numbering_is_continuous(basic_model, tmp_path):
    basic_model["sections"] = [
        {"level": 1, "text": "1 引言"},
        {"type": "equation", "latex": "a = b"},
        {"type": "equation", "latex": "c = d"},
    ]
    _, path = db.build_from_model(basic_model, str(tmp_path / "e.docx"))
    doc = Document(path)
    eqs = [t for t in doc.tables if len(t.columns) == 3]
    numbers = [t.rows[0].cells[2].text.strip() for t in eqs]
    assert numbers == ["（1）", "（2）"]


# --------------------------------------------------------------- 页眉页脚

def test_header_footer_content(built_doc):
    sec = built_doc.sections[0]
    header = sec.header.paragraphs[0]
    assert header.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert "测试标题" in header.text
    assert " · " in header.text
    size = next(r.font.size.pt for r in header.runs if r.font.size)
    assert size == 8.5

    footer = sec.footer.paragraphs[0]
    instrs = " ".join(el.text or "" for el in footer._p.iter(w("instrText")))
    assert "PAGE" in instrs and "NUMPAGES" in instrs
    assert "第" in footer.text and "页" in footer.text
    assert footer.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert next(r.font.size.pt for r in footer.runs if r.font.size) == 8


# --------------------------------------------------------------- 参考文献

def test_references_hanging_indent(built_doc):
    refs = [p for p in built_doc.paragraphs if p.text.startswith("[1]")]
    assert refs
    ind = refs[0]._p.find(w("pPr")).find(w("ind"))
    assert ind.get(w("hanging")) == "420"
    assert ind.get(w("left")) == "420"
    assert next(r.font.size.pt for r in refs[0].runs if r.font.size) == 9


def test_references_get_numbers_when_missing(basic_model, tmp_path):
    basic_model["references"] = ["甲. 题名[J]. 刊名, 2020, 1(1): 1-2.",
                                 "乙. 题名[M]. 北京: 出版社, 2020."]
    _, path = db.build_from_model(basic_model, str(tmp_path / "r.docx"))
    doc = Document(path)
    refs = [p.text for p in doc.paragraphs if p.text.strip().startswith("[")]
    assert refs[0].startswith("[1]") and refs[1].startswith("[2]")


# --------------------------------------------------------------- 目录 / spec

def test_toc_field(built_doc):
    body = built_doc.element.body
    instrs = [el.text or "" for el in body.iter(w("instrText"))]
    toc = [x for x in instrs if "TOC" in x]
    assert toc, "没有 TOC 域"
    assert '"1-2"' in toc[0]


def test_spec_override_changes_palette(tmp_path):
    builder = db.PaperBuilder({"colors": {"navy": "123456", "blue": "654321"}})
    builder.setup()
    assert builder._c("navy") == "123456"
    style = builder.doc.styles["Heading 1"]
    assert style.element.find(w("rPr")).find(w("color")).get(w("val")) == "123456"


def test_aligns_accept_strings(basic_model, tmp_path):
    basic_model["sections"] = [
        {"level": 1, "text": "1 引言"},
        {"type": "table", "caption": "对齐表", "headers": ["A", "B"],
         "rows": [["1", "2"]], "aligns": ["right", "left"]},
    ]
    _, path = db.build_from_model(basic_model, str(tmp_path / "a.docx"))
    doc = Document(path)
    cell = doc.tables[-1].rows[1].cells[0]
    assert cell.paragraphs[0].alignment is not None


def test_missing_figure_raises(basic_model, tmp_path):
    basic_model["sections"] = [
        {"level": 1, "text": "1 引言"},
        {"type": "figure", "path": str(tmp_path / "nope.png"), "caption": "没有这张图"},
    ]
    with pytest.raises(FileNotFoundError):
        db.build_from_model(basic_model, str(tmp_path / "x.docx"))


def test_build_from_demo_model(tmp_path):
    """样例内容模型必须能直接构建（防止样例与代码脱节）。"""
    model_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                              "examples", "demo", "content_model.json")
    builder, path = db.build_from_model(
        demo_model(), str(tmp_path / "demo.docx"),
        base_dir=os.path.dirname(model_path))
    assert os.path.exists(path)
    assert builder.tbl_no >= 1 and builder.fig_no >= 1 and builder.eq_no >= 1
    assert builder.h1_no >= 3


def test_paragraph_helpers_do_not_leak_spec_state():
    a = db.PaperBuilder()
    b = db.PaperBuilder({"colors": {"navy": "000000"}})
    assert a.spec["colors"]["navy"] == "1F3A5F"
    assert b.spec["colors"]["navy"] == "000000"


def test_hanging_cm_parameter(tmp_path):
    """负缩进必须写成 w:hanging，且单位是 twips（历史 bug 回归测试）。"""
    builder = db.PaperBuilder()
    builder.setup()
    builder.add_references(["甲. 题名[J]. 刊名, 2020."])
    p = builder.doc.paragraphs[-1]
    ind = p._p.find(w("pPr")).find(w("ind"))
    assert ind.get(w("hanging")) == "420"
    assert int(ind.get(w("hanging"))) < 6320  # 不能是 EMU 数量级


def test_caption_pt_and_color(built_doc):
    # 题注 = 表号 + 全角空格 + 题名；正文里"表1 给出…"不是题注
    cap = next(p for p in built_doc.paragraphs
               if p.text.startswith("\u88681\u3000"))
    run = cap.runs[0]
    assert run.font.size == Pt(9.5)
    assert str(run.font.color.rgb) == "1F3A5F"
    assert run.bold is True


# ---------------------------------------------------------------- 国标条款

def test_format_number_blank_semantics():
    """5.5.4：空白 = 无此项或未测量；“—” = 测量过而未发现；0 = 实为零。"""
    assert db.format_number(None) == ""
    assert db.format_number(None, 2) == ""
    assert db.format_number("") == ""
    assert db.format_number("—") == "—"
    assert db.format_number(0) == "0"


def test_caption_uses_fullwidth_space(built):
    """5.4.3 / 5.5.3：图号、表号与题名之间空 1 个汉字。"""
    b = built[0]
    _, fig_text = b.add_figure_caption("测试图题")
    assert "\u3000" in fig_text and "  " not in fig_text
    _, tbl_text = b.add_table_caption("测试表题")
    assert "\u3000" in tbl_text


def test_note_gets_period(built):
    """5.4.3 / 5.5.3：图注、表注末尾应加“。”。"""
    b = built[0]
    p = b.add_note("数据为模拟结果")
    assert p.text.strip().endswith("。")


def test_figure_width_clamped_to_text_block(built):
    """图宽超过版心时自动收窄并警告。"""
    b = built[0]
    b.add_figure(DEMO_FIGURE, caption="超宽测试", width_cm=20.0)
    assert any("版心" in w for w in b.warnings)


def test_figure_font_shrinks_when_inserted_narrow(built, tmp_path):
    """155 mm 设计的图插入半栏 → 有效字号低于 6 pt 应告警。"""
    import figures
    spec = {"kind": "line", "width_mm": 155, "dpi": 300,
            "series": [{"name": "A", "values": [1, 2, 3]}], "y_range": [0, 4]}
    png = str(tmp_path / "wide.png")
    figures.render(spec, png)
    b = built[0]
    b.add_figure(png, caption="缩太小", width_cm=7.5)
    assert any("pt" in w and "低于" in w for w in b.warnings)


# ---------------------------------------------------------------- 前置部分顺序

def _front_titles(doc):
    out = []
    for para in doc.paragraphs:
        t = para.text.replace("\u3000", "").strip()
        if t in ("摘要", "目录"):
            out.append(t)
        elif t.startswith("Abstract"):
            out.append("Abstract")
    return out


def test_default_front_order_is_abstract_before_toc(basic_model, tmp_path):
    """GB/T 7713.1—2025 附录 A 表 A.1：摘要页在目次页之前。"""
    model = dict(basic_model)
    model["abstract_en"] = {"paragraphs": ["English abstract for ordering test."],
                            "keywords": ["kw"]}
    _, path = db.build_from_model(model, str(tmp_path / "order.docx"))
    seq = _front_titles(Document(path))
    assert seq[:3] == ["摘要", "Abstract", "目录"], seq


def test_front_order_can_put_toc_first(basic_model, tmp_path):
    """少数学校模板要求目录在前：front_order = "toc-first"。"""
    model = dict(basic_model)
    model["abstract_en"] = {"paragraphs": ["English abstract."], "keywords": ["kw"]}
    model["front_order"] = "toc-first"
    _, path = db.build_from_model(model, str(tmp_path / "tocfirst.docx"))
    seq = _front_titles(Document(path))
    assert seq[:3] == ["目录", "摘要", "Abstract"], seq


def test_english_abstract_starts_on_new_page(basic_model, tmp_path):
    """5.2.9：英文摘要宜另起一页，置于中文摘要页之后。"""
    model = dict(basic_model)
    model["abstract_en"] = {"paragraphs": ["English abstract."], "keywords": ["kw"]}
    _, path = db.build_from_model(model, str(tmp_path / "en.docx"))
    doc = Document(path)
    abs_para = next(p for p in doc.paragraphs if p.text.strip().startswith("Abstract"))
    pPr = abs_para._p.find(w("pPr"))
    assert pPr is not None and pPr.find(w("pageBreakBefore")) is not None


# ---------------------------------------------------------------- 三级标题 / 目录层级

def test_heading_level_3_supported(basic_model, tmp_path):
    """三级标题：Heading 3 样式、12 pt 加粗、大纲级别 2。"""
    model = dict(basic_model)
    model["sections"] = [
        {"level": 1, "text": "1 引言"},
        {"level": 2, "text": "1.1 小节"},
        {"level": 3, "text": "1.1.1 更小的节"},
        {"type": "p", "text": "正文。"},
    ]
    _, path = db.build_from_model(model, str(tmp_path / "h3.docx"))
    doc = Document(path)
    h3 = next(p for p in doc.paragraphs if p.text.strip() == "1.1.1 更小的节")
    assert h3.style.name == "Heading 3"
    assert h3.runs[0].font.size == Pt(12)
    pPr = h3.style.element.find(w("pPr"))
    assert pPr is not None and pPr.find(w("outlineLvl")) is not None


def test_toc_default_lists_two_levels(basic_model, tmp_path):
    """目录默认列到二级（TOC \\o "1-2"）：只列一级几乎无法导航。"""
    _, path = db.build_from_model(basic_model, str(tmp_path / "toc2.docx"))
    doc = Document(path)
    instr = [t.text for t in doc.element.body.iter(qn("w:instrText")) if "TOC" in t.text]
    assert instr and '"1-2"' in instr[0], instr


def test_toc_levels_configurable_to_three(basic_model, tmp_path):
    model = dict(basic_model)
    model["toc_levels"] = "1-3"
    _, path = db.build_from_model(model, str(tmp_path / "toc3.docx"))
    doc = Document(path)
    instr = [t.text for t in doc.element.body.iter(qn("w:instrText")) if "TOC" in t.text]
    assert instr and '"1-3"' in instr[0], instr
